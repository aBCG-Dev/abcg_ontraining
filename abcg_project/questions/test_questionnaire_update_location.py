import io
import openpyxl
from django.test import TestCase, Client
from django.urls import reverse
from django.contrib.auth.models import User
from django.utils import timezone
from questions.models import (
    Participant, IneligibleIndividual, TptIndividual,
    UserProfile, RolePermission, Symptom, RiskFactor
)

class QuestionnaireUpdateLocationPreservationTestCase(TestCase):
    def setUp(self):
        self.client = Client()

        # Seed basic symptoms and risk factors if needed
        Symptom.objects.get_or_create(code="fever", defaults={"name": "Fever", "is_active": True})
        Symptom.objects.get_or_create(code="cough", defaults={"name": "Cough > 2 weeks", "is_active": True})
        RiskFactor.objects.get_or_create(code="diabetes", defaults={"name": "Diabetes", "is_active": True})
        RiskFactor.objects.get_or_create(code="smoker", defaults={"name": "Tobacco/smoker", "is_active": True})

        # 1. Global Admin User
        self.global_admin = User.objects.create_user(
            username="global_admin", password="password123", email="global@example.com"
        )
        self.global_admin_profile = self.global_admin.profile
        self.global_admin_profile.full_name = "Global Super Admin"
        self.global_admin_profile.role = "Super Admin"
        self.global_admin_profile.state = "GLOBAL"
        self.global_admin_profile.district = "ALL DISTRICTS"
        self.global_admin_profile.tb_unit = "ALL TB UNITS"
        self.global_admin_profile.save()

        # 2. State-level Admin User (Maharashtra)
        self.state_admin = User.objects.create_user(
            username="state_admin", password="password123", email="state@example.com"
        )
        self.state_admin_profile = self.state_admin.profile
        self.state_admin_profile.full_name = "Maharashtra State Admin"
        self.state_admin_profile.role = "Admin"
        self.state_admin_profile.state = "Maharashtra"
        self.state_admin_profile.district = "ALL DISTRICTS"
        self.state_admin_profile.tb_unit = "ALL TB UNITS"
        self.state_admin_profile.save()

        # 3. District-level Admin User (Chennai, Tamil Nadu)
        self.district_admin = User.objects.create_user(
            username="district_admin", password="password123", email="district@example.com"
        )
        self.district_admin_profile = self.district_admin.profile
        self.district_admin_profile.full_name = "Chennai District Admin"
        self.district_admin_profile.role = "Admin"
        self.district_admin_profile.state = "Tamil Nadu"
        self.district_admin_profile.district = "Chennai"
        self.district_admin_profile.tb_unit = "ALL TB UNITS"
        self.district_admin_profile.save()

        # 4. TB Unit User (Chennai Central TU, Chennai, Tamil Nadu)
        self.tu_nurse = User.objects.create_user(
            username="tu_nurse", password="password123", email="nurse@example.com"
        )
        self.tu_nurse_profile = self.tu_nurse.profile
        self.tu_nurse_profile.full_name = "Nurse Chennai Central"
        self.tu_nurse_profile.role = "Project Nurse"
        self.tu_nurse_profile.state = "Tamil Nadu"
        self.tu_nurse_profile.district = "Chennai"
        self.tu_nurse_profile.tb_unit = "Chennai Central TU"
        self.tu_nurse_profile.save()

        # 5. Local Nurse at Tiruvallur TU (Tiruvallur, Tamil Nadu)
        self.tiruvallur_nurse = User.objects.create_user(
            username="tiruvallur_nurse", password="password123", email="tirunurse@example.com"
        )
        self.tiruvallur_nurse_profile = self.tiruvallur_nurse.profile
        self.tiruvallur_nurse_profile.full_name = "Nurse Tiruvallur"
        self.tiruvallur_nurse_profile.role = "Project Nurse"
        self.tiruvallur_nurse_profile.state = "Tamil Nadu"
        self.tiruvallur_nurse_profile.district = "Tiruvallur"
        self.tiruvallur_nurse_profile.tb_unit = "Tiruvallur TU"
        self.tiruvallur_nurse_profile.save()

        # Grant permissions for Admin and Project Nurse in RolePermission table if checked
        for role in ["Admin", "Project Nurse"]:
            for action in ["View", "Create", "Edit", "Export"]:
                RolePermission.objects.get_or_create(
                    role=role, module="Participant Registration", action=action,
                    defaults={"allowed": True}
                )

        # Base payload helper
        self.valid_payload = {
            "first_name": "Ramesh",
            "last_name": "Kumar",
            "father_husband_name": "Suresh Kumar",
            "age": 42,
            "dob": "1984-06-15",
            "gender": "Male",
            "primary_phone": "9876543210",
            "secondary_phone": "9876543211",
            "address": "12 Gandhi Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "taluka_block": "Avadi Taluka",
            "village": "Avadi",
            "pincode": "600054",
            "demographic_area": "Urban",
            "marital_status": "Married",
            "occupation": "Shopkeeper",
            "socioeconomic_status": "APL",
            "symptoms": ["Fever"],
            "risk_factors": ["Tobacco/smoker"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"]
        }

    def _create_sample_participant(self, state="Tamil Nadu", district="Tiruvallur", tb_unit="Tiruvallur TU"):
        return Participant.objects.create(
            study_id="ABCG-TN-TIR-2026-123456",
            screening_id="SCR-260101-1001",
            nikshay_id="NIK12345678",
            first_name="Ramesh",
            last_name="Kumar",
            full_name="Ramesh Kumar",
            contact_number="9876543210",
            age=42,
            gender="Male",
            state=state,
            district=district,
            tb_unit=tb_unit,
            address="12 Gandhi Road",
            date_enroll=timezone.localdate(),
            created_by=self.tiruvallur_nurse,
            eligible=True,
            classification="Case"
        )

    # 1. A Global administrator updates an existing participant's questionnaire
    def test_01_global_admin_updates_participant_preserves_location(self):
        p = self._create_sample_participant()
        self.client.force_login(self.global_admin)

        # GET: Verify form renders participant's actual site location, NEVER Global admin's GLOBAL scope
        get_resp = self.client.get(reverse("questions:registration_edit", args=[p.pk]))
        self.assertEqual(get_resp.status_code, 200)
        self.assertEqual(get_resp.context["site"]["state"], "Tamil Nadu", "GET site.state must be participant's original location")
        self.assertEqual(get_resp.context["site"]["district"], "Tiruvallur")
        self.assertEqual(get_resp.context["site"]["tb_unit"], "Tiruvallur TU")

        payload = dict(self.valid_payload)
        payload["first_name"] = "RameshUpdatedByGlobal"
        payload["state"] = "GLOBAL"
        payload["district"] = "ALL DISTRICTS"
        payload["tb_unit"] = "ALL TB UNITS"

        response = self.client.post(reverse("questions:registration_edit", args=[p.pk]), data=payload)
        self.assertRedirects(response, reverse("questions:home"), fetch_redirect_response=False)

        p.refresh_from_db()
        self.assertEqual(p.first_name, "RameshUpdatedByGlobal")
        self.assertEqual(p.state, "Tamil Nadu", "Participant state must not be overwritten by Global admin scope")
        self.assertEqual(p.district, "Tiruvallur", "Participant district must not be overwritten by Global admin scope")
        self.assertEqual(p.tb_unit, "Tiruvallur TU", "Participant tb_unit must not be overwritten by Global admin scope")

        # Healing test: If a participant was corrupted to GLOBAL, editing restores their verified true location
        p.state = "GLOBAL"
        p.district = "ALL DISTRICTS"
        p.tb_unit = "ALL TB UNITS"
        p.save()

        get_corrupted = self.client.get(reverse("questions:registration_edit", args=[p.pk]))
        self.assertEqual(get_corrupted.status_code, 200)
        self.assertEqual(get_corrupted.context["site"]["state"], "Tamil Nadu", "Must heal corrupted GLOBAL back to Tamil Nadu from Study ID")
        self.assertEqual(get_corrupted.context["site"]["district"], "Tiruvallur")
        self.assertEqual(get_corrupted.context["site"]["tb_unit"], "Tiruvallur TU")

        # Saving also heals the corrupted database record
        post_heal = self.client.post(reverse("questions:registration_edit", args=[p.pk]), data=payload)
        self.assertRedirects(post_heal, reverse("questions:home"), fetch_redirect_response=False)
        p.refresh_from_db()
        self.assertEqual(p.state, "Tamil Nadu", "Healed record in database must be Tamil Nadu")
        self.assertEqual(p.district, "Tiruvallur")
        self.assertEqual(p.tb_unit, "Tiruvallur TU")

    # 2. A state-level administrator updates a participant's questionnaire
    def test_02_state_admin_updates_participant_preserves_location(self):
        p = self._create_sample_participant()
        self.client.force_login(self.state_admin)

        payload = dict(self.valid_payload)
        payload["first_name"] = "RameshUpdatedByStateAdmin"
        payload["state"] = "Maharashtra"
        payload["district"] = "ALL DISTRICTS"
        payload["tb_unit"] = "ALL TB UNITS"

        response = self.client.post(reverse("questions:registration_edit", args=[p.pk]), data=payload)
        self.assertRedirects(response, reverse("questions:home"), fetch_redirect_response=False)

        p.refresh_from_db()
        self.assertEqual(p.first_name, "RameshUpdatedByStateAdmin")
        self.assertEqual(p.state, "Tamil Nadu", "Participant state must not be changed by State admin")
        self.assertEqual(p.district, "Tiruvallur", "Participant district must not be changed by State admin")
        self.assertEqual(p.tb_unit, "Tiruvallur TU", "Participant tb_unit must not be changed by State admin")

    # 3. A district-level administrator updates a participant's questionnaire
    def test_03_district_admin_updates_participant_preserves_location(self):
        p = self._create_sample_participant()
        self.client.force_login(self.district_admin)

        payload = dict(self.valid_payload)
        payload["first_name"] = "RameshUpdatedByDistrictAdmin"
        payload["state"] = "Tamil Nadu"
        payload["district"] = "Chennai"
        payload["tb_unit"] = "ALL TB UNITS"

        response = self.client.post(reverse("questions:registration_edit", args=[p.pk]), data=payload)
        self.assertRedirects(response, reverse("questions:home"), fetch_redirect_response=False)

        p.refresh_from_db()
        self.assertEqual(p.first_name, "RameshUpdatedByDistrictAdmin")
        self.assertEqual(p.state, "Tamil Nadu")
        self.assertEqual(p.district, "Tiruvallur", "Participant district must be preserved as Tiruvallur")
        self.assertEqual(p.tb_unit, "Tiruvallur TU", "Participant tb_unit must be preserved as Tiruvallur TU")

    # 4. A TB Unit-level user updates a participant's questionnaire
    def test_04_tb_unit_user_updates_participant_preserves_location(self):
        p = self._create_sample_participant()
        self.client.force_login(self.tu_nurse)

        payload = dict(self.valid_payload)
        payload["first_name"] = "RameshUpdatedByTuNurse"
        payload["state"] = "Tamil Nadu"
        payload["district"] = "Chennai"
        payload["tb_unit"] = "Chennai Central TU"

        response = self.client.post(reverse("questions:registration_edit", args=[p.pk]), data=payload)
        self.assertRedirects(response, reverse("questions:home"), fetch_redirect_response=False)

        p.refresh_from_db()
        self.assertEqual(p.first_name, "RameshUpdatedByTuNurse")
        self.assertEqual(p.state, "Tamil Nadu")
        self.assertEqual(p.district, "Tiruvallur")
        self.assertEqual(p.tb_unit, "Tiruvallur TU")

    # 5. An ineligible participant becomes eligible and receives a Study ID
    def test_05_ineligible_becomes_eligible_derives_study_id_from_actual_location(self):
        # Create an ineligible participant with SCR- id
        ineligible = IneligibleIndividual.objects.create(
            study_id="SCR-260101-9999",
            screening_id="SCR-260101-9999",
            nikshay_id="NIK99999999",
            first_name="Anand",
            last_name="Verma",
            full_name="Anand Verma",
            contact_number="9876500000",
            age=30,
            gender="Male",
            state="Tamil Nadu",
            district="Tiruvallur",
            tb_unit="Tiruvallur TU",
            date_enroll=timezone.localdate(),
            eligible=False,
            classification="Ineligible"
        )

        # Global Admin updates the ineligible participant into an eligible Case
        self.client.force_login(self.global_admin)

        payload = dict(self.valid_payload)
        payload["first_name"] = "Anand"
        payload["last_name"] = "Verma"
        payload["primary_phone"] = "9876500000"
        payload["past_tb"] = "Yes"
        payload["symptoms"] = ["Fever", "Cough > 2 weeks"]
        payload["eligible_bcg_campaign"] = "Yes"
        payload["bcg_eligibility_criteria"] = ["Previously treated for TB disease"]

        response = self.client.post(reverse("questions:registration_edit", args=[ineligible.pk]), data=payload)
        self.assertRedirects(response, reverse("questions:home"), fetch_redirect_response=False)

        # Verify old IneligibleIndividual deleted and new Participant created
        self.assertFalse(IneligibleIndividual.objects.filter(screening_id="SCR-260101-9999").exists())
        promoted = Participant.objects.get(screening_id="SCR-260101-9999")

        # Location must remain Tamil Nadu, Tiruvallur, Tiruvallur TU
        self.assertEqual(promoted.state, "Tamil Nadu")
        self.assertEqual(promoted.district, "Tiruvallur")
        self.assertEqual(promoted.tb_unit, "Tiruvallur TU")

        # Study ID must start with ABCG-TN-TIR- and NOT ABCG-GL- or ABCG-ALL-
        self.assertTrue(promoted.study_id.startswith("ABCG-TN-TIR-"), 
                        f"Study ID '{promoted.study_id}' must derive from participant's location (TN-TIR), not Global profile")

    # 6. A new participant registration is submitted
    def test_06_new_registration_submitted_uses_correct_location_and_study_id(self):
        self.client.force_login(self.tiruvallur_nurse)

        payload = dict(self.valid_payload)
        payload["first_name"] = "NewParticipant"
        payload["last_name"] = "Tiruvallur"
        payload["primary_phone"] = "9111222333"

        response = self.client.post(reverse("questions:registration"), data=payload)
        self.assertRedirects(response, reverse("questions:home"), fetch_redirect_response=False)

        new_p = Participant.objects.get(first_name="NewParticipant", last_name="Tiruvallur")
        self.assertEqual(new_p.state, "Tamil Nadu")
        self.assertEqual(new_p.district, "Tiruvallur")
        self.assertEqual(new_p.tb_unit, "Tiruvallur TU")
        self.assertTrue(new_p.study_id.startswith("ABCG-TN-TIR-2026-"))

    # 7. An existing questionnaire is edited without changing any location fields
    def test_07_edit_questionnaire_without_location_change_preserves_location(self):
        p = self._create_sample_participant()
        self.client.force_login(self.tiruvallur_nurse)

        payload = dict(self.valid_payload)
        payload["first_name"] = "Ramesh"
        payload["last_name"] = "Kumar"
        payload["occupation"] = "Senior Engineer"
        payload["marital_status"] = "Divorced"
        payload["symptoms"] = ["Cough > 2 weeks"]

        response = self.client.post(reverse("questions:registration_edit", args=[p.pk]), data=payload)
        self.assertRedirects(response, reverse("questions:home"), fetch_redirect_response=False)

        p.refresh_from_db()
        self.assertEqual(p.occupation, "Senior Engineer")
        self.assertEqual(p.marital_status, "Divorced")
        self.assertEqual(p.state, "Tamil Nadu")
        self.assertEqual(p.district, "Tiruvallur")
        self.assertEqual(p.tb_unit, "Tiruvallur TU")

    # 8. The updated records are exported to Excel
    def test_08_exported_excel_retains_correct_locations(self):
        p = self._create_sample_participant()
        self.client.force_login(self.global_admin)

        response = self.client.get(reverse("questions:data_export") + "?format=excel")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        # Parse Excel file
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn("Master Registry", wb.sheetnames)
        ws = wb["Master Registry"]

        # Find column indices for State, District, TB Unit
        headers = [cell.value for cell in ws[1]]
        state_idx = headers.index("State") + 1
        dist_idx = headers.index("District") + 1
        tu_idx = headers.index("TB Unit") + 1

        # Check rows
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertGreaterEqual(len(rows), 1)

        found_sample = False
        for row in rows:
            row_state = row[state_idx - 1]
            row_dist = row[dist_idx - 1]
            row_tu = row[tu_idx - 1]

            # Ensure none of the rows are corrupted with GLOBAL / ALL
            self.assertNotEqual(row_state, "GLOBAL")
            self.assertNotEqual(row_dist, "ALL DISTRICTS")
            self.assertNotEqual(row_tu, "ALL TB UNITS")

            if row_state == "Tamil Nadu" and row_dist == "Tiruvallur" and row_tu == "Tiruvallur TU":
                found_sample = True

        self.assertTrue(found_sample, "Sample participant with Tamil Nadu / Tiruvallur / Tiruvallur TU must be present in exported Excel")
