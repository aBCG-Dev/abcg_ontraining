import io
import datetime
import openpyxl
from django.test import TestCase, Client
from django.urls import reverse
from django.utils import timezone
from django.contrib.auth.models import User
from questions.models import (
    Participant, IneligibleIndividual, TptIndividual,
    UserProfile, RolePermission, Symptom, RiskFactor
)

class EnrollmentDatePreservationTestCase(TestCase):
    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(username="admin_user", password="password123")
        self.user.profile.role = "Super Admin"
        self.user.profile.save()
        self.client.force_login(self.user)

        # Seed basic symptoms and risk factors
        Symptom.objects.get_or_create(code="fever", defaults={"name": "Fever", "is_active": True})
        RiskFactor.objects.get_or_create(code="diabetes", defaults={"name": "Diabetes", "is_active": True})

        self.historical_date = datetime.date(2025, 11, 20)
        self.base_payload = {
            "first_name": "Ramesh",
            "last_name": "Kumar",
            "father_husband_name": "Suresh Kumar",
            "age": 42,
            "dob": "1984-06-15",
            "gender": "Male",
            "primary_phone": "9876543210",
            "secondary_phone": "9876543211",
            "address": "12 Gandhi Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "taluka_block": "Avadi Taluka",
            "village": "Avadi",
            "pincode": "600054",
            "demographic_area": "Urban",
            "marital_status": "Married",
            "occupation": "Shopkeeper",
            "socioeconomic_status": "APL",
            "symptoms": ["Fever"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"]
        }

    def test_01_update_participant_preserves_enrollment_date_and_records_updated_at(self):
        p = Participant.objects.create(
            study_id="ABCG-TN-TIR-2025-111222",
            screening_id="SCR-251120-1001",
            first_name="Ramesh",
            last_name="Kumar",
            full_name="Ramesh Kumar",
            contact_number="9876543210",
            age=42,
            gender="Male",
            state="Tamil Nadu",
            district="Tiruvallur",
            tb_unit="Tiruvallur TU",
            date_enroll=self.historical_date,
            eligible=True,
            classification="Case",
            created_by=self.user
        )
        original_created_at = p.created_at

        payload = self.base_payload.copy()
        payload["occupation"] = "Senior Manager"
        response = self.client.post(reverse("questions:registration_edit", kwargs={"pk": p.pk}), data=payload)
        self.assertRedirects(response, reverse("questions:home"), fetch_redirect_response=False)

        p.refresh_from_db()
        self.assertEqual(p.occupation, "Senior Manager")
        # Enrollment date must NOT change to current date
        self.assertEqual(p.date_enroll, self.historical_date, f"Expected {self.historical_date}, got {p.date_enroll}")
        # updated_at must be populated and recent
        self.assertIsNotNone(p.updated_at)
        self.assertGreaterEqual(p.updated_at, original_created_at)

    def test_02_update_screening_draft_preserves_enrollment_date(self):
        # Participant created from screening register with SCR- ID and historical date
        p = Participant.objects.create(
            study_id="SCR-251120-9999",
            screening_id="SCR-251120-9999",
            first_name="ScreeningDraft",
            last_name="User",
            full_name="ScreeningDraft User",
            contact_number="9876543210",
            age=42,
            gender="Male",
            state="Tamil Nadu",
            district="Tiruvallur",
            tb_unit="Tiruvallur TU",
            date_enroll=self.historical_date,
            eligible=True,
            classification="Pending",
            created_by=self.user
        )

        payload = self.base_payload.copy()
        payload["first_name"] = "ScreeningConfirmed"
        response = self.client.post(reverse("questions:registration_edit", kwargs={"pk": p.pk}), data=payload)
        self.assertRedirects(response, reverse("questions:home"), fetch_redirect_response=False)

        p.refresh_from_db()
        self.assertEqual(p.first_name, "ScreeningConfirmed")
        self.assertEqual(p.date_enroll, self.historical_date, f"Draft promotion must preserve date_enroll {self.historical_date}, got {p.date_enroll}")

    def test_03_cohort_transition_ineligible_to_participant_preserves_date(self):
        # Ineligible individual transitioning to eligible Participant
        ineligible = IneligibleIndividual.objects.create(
            study_id="SCR-251120-5555",
            screening_id="SCR-251120-5555",
            first_name="IneligibleUser",
            last_name="Test",
            full_name="IneligibleUser Test",
            contact_number="9876543210",
            age=42,
            gender="Male",
            state="Tamil Nadu",
            district="Tiruvallur",
            tb_unit="Tiruvallur TU",
            date_enroll=self.historical_date,
            eligible=False,
            classification="Not Eligible",
            created_by=self.user
        )

        payload = self.base_payload.copy()
        payload["first_name"] = "NowEligible"
        response = self.client.post(reverse("questions:registration_edit", kwargs={"pk": ineligible.pk}) + "?cohort=ineligible", data=payload)
        self.assertRedirects(response, reverse("questions:home"), fetch_redirect_response=False)

        promoted = Participant.objects.get(screening_id="SCR-251120-5555")
        self.assertEqual(promoted.first_name, "NowEligible")
        self.assertEqual(promoted.date_enroll, self.historical_date, f"Cohort transition must preserve date_enroll {self.historical_date}, got {promoted.date_enroll}")
        self.assertIsNotNone(promoted.updated_at)

    def test_04_partial_update_preserves_unmodified_fields(self):
        p = Participant.objects.create(
            study_id="ABCG-TN-TIR-2025-333444",
            screening_id="SCR-251120-3334",
            first_name="OriginalFirst",
            last_name="OriginalLast",
            father_husband_name="OriginalFather",
            contact_number="9876543210",
            address="Original Address 123",
            age=45,
            gender="Male",
            state="Tamil Nadu",
            district="Tiruvallur",
            tb_unit="Tiruvallur TU",
            facility="Avadi PHC",
            public_phi="Avadi PHC",
            village="Avadi Village",
            pincode="600054",
            date_enroll=self.historical_date,
            eligible=True,
            classification="Case",
            created_by=self.user
        )

        # Submit update with only first_name and address changed; omitted fields must not be wiped
        payload = {
            "first_name": "ModifiedFirst",
            "address": "New Address 456",
            "age": 45,
            "gender": "Male",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"],
            "risk_factors": ["Diabetes"],
            "symptoms": ["Fever"]
        }
        response = self.client.post(reverse("questions:registration_edit", kwargs={"pk": p.pk}), data=payload)
        self.assertRedirects(response, reverse("questions:home"), fetch_redirect_response=False)

        p.refresh_from_db()
        self.assertEqual(p.first_name, "ModifiedFirst")
        self.assertEqual(p.address, "New Address 456")
        # Unmodified fields must be preserved
        self.assertEqual(p.last_name, "OriginalLast")
        self.assertEqual(p.father_husband_name, "OriginalFather")
        self.assertEqual(p.contact_number, "9876543210")
        self.assertEqual(p.date_enroll, self.historical_date)

    def test_05_excel_export_displays_original_enrollment_date(self):
        p = Participant.objects.create(
            study_id="ABCG-TN-TIR-2025-777888",
            screening_id="SCR-251120-7778",
            first_name="ExcelTest",
            last_name="User",
            full_name="ExcelTest User",
            contact_number="9876543210",
            age=38,
            gender="Female",
            state="Tamil Nadu",
            district="Tiruvallur",
            tb_unit="Tiruvallur TU",
            date_enroll=self.historical_date,
            eligible=True,
            classification="Control",
            created_by=self.user
        )

        export_url = reverse("questions:data_export") + "?format=excel"
        response = self.client.get(export_url)
        self.assertEqual(response.status_code, 200)

        # Parse Excel output with openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active

        headers = [cell.value for cell in ws[1]]
        self.assertIn("Date of Enrollment", headers)
        self.assertIn("Last Updated Date", headers)

        date_enroll_col = headers.index("Date of Enrollment") + 1
        study_id_col = headers.index("Study ID") + 1

        row_found = False
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=study_id_col).value == p.study_id:
                row_found = True
                val = ws.cell(row=row, column=date_enroll_col).value
                # Value should match the historical date
                if isinstance(val, (datetime.date, datetime.datetime)):
                    val_date = val.date() if isinstance(val, datetime.datetime) else val
                else:
                    val_date = datetime.datetime.strptime(str(val)[:10], "%Y-%m-%d").date()
                self.assertEqual(val_date, self.historical_date, f"Excel Date of Enrollment must be {self.historical_date}, got {val}")
                break

        self.assertTrue(row_found, f"Participant {p.study_id} should appear in Excel export")

    def test_06_new_registration_workflow_sets_today_date(self):
        payload = self.base_payload.copy()
        payload["first_name"] = "BrandNewParticipant"
        payload["primary_phone"] = "9111122222"

        response = self.client.post(reverse("questions:registration"), data=payload)
        self.assertRedirects(response, reverse("questions:home"), fetch_redirect_response=False)

        new_p = Participant.objects.get(first_name="BrandNewParticipant")
        self.assertEqual(new_p.date_enroll, timezone.localdate(), "New participant must have today's enrollment date")
        self.assertIsNotNone(new_p.created_at)
        self.assertIsNotNone(new_p.updated_at)
