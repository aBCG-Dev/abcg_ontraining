import json
from django.test import TestCase
from django.urls import reverse, resolve
from django.contrib.auth.models import User
from django.utils import timezone
from questions.models import Participant
from questions import views

class QuestionsRoutingAndViewsTestCase(TestCase):
    def setUp(self):
        # Create a test user and their profile is created automatically by signals
        self.username = "teststaff"
        self.password = "password123"
        self.user = User.objects.create_user(
            username=self.username,
            email="teststaff@example.com",
            password=self.password
        )
        self.user.first_name = "Test"
        self.user.last_name = "Staff"
        self.user.save()
        
    def test_home_url_resolves(self):
        url = reverse("questions:home")
        self.assertEqual(resolve(url).func, views.home)
        
    def test_home_page_renders_login_for_anonymous(self):
        response = self.client.get(reverse("questions:home"))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context["is_login_page"])
        self.assertTemplateUsed(response, "questions/home.html")
        
    def test_login_post_redirects_to_home(self):
        response = self.client.post(reverse("questions:home"), {
            "username": self.username,
            "password": self.password
        })
        self.assertRedirects(response, reverse("questions:home"))
        
    def test_login_post_redirects_to_dashboard_for_doctor(self):
        profile = self.user.profile
        profile.role = "Doctor"
        profile.save()
        response = self.client.post(reverse("questions:home"), {
            "username": self.username,
            "password": self.password
        })
        self.assertRedirects(response, reverse("questions:dashboard_home"))
        
    def test_home_page_renders_dashboard_for_authenticated(self):
        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:home"))
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context["is_login_page"])
        self.assertEqual(response.context["site"]["state"], "Tamil Nadu")
        
    def test_home_page_redirects_for_authenticated_doctor(self):
        profile = self.user.profile
        profile.role = "Doctor"
        profile.save()
        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:home"))
        self.assertRedirects(response, reverse("questions:dashboard_home"))
        
    def test_registration_page_redirects_anonymous(self):
        response = self.client.get(reverse("questions:registration"))
        self.assertRedirects(response, reverse("questions:home"))
        
    def test_registration_page_get_authenticated(self):
        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:registration"))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "questions/registration.html")
        
    def test_registration_page_post_creates_participant(self):
        self.client.login(username=self.username, password=self.password)
        # Verify initial database count
        from questions.models import Participant
        initial_count = Participant.objects.count()
        
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Muni",
            "last_name": "Ramana",
            "father_husband_name": "Father Name",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543210",
            "secondary_phone": "9123456789",
            "secondary_phone_1": "9876543211",
            "secondary_phone_2": "9876543212",
            "secondary_phone_3": "9876543213",
            "address": "123 Temple Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "taluka_block": "Avadi Taluka",
            "landmark": "Near Temple",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            "contact_person_name": "Contact Name",
            "contact_person_phone": "9876543219",
            "contact_person_address": "456 Main St",
            "informant_name": "Informant Name",
            "informant_designation": "Staff Nurse",
            "symptoms": ["Fever", "Weight loss"],
            "risk_factors": ["Diabetes", "Tobacco/smoker"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"]
        })
        self.assertRedirects(response, reverse("questions:home"))
        
        # Verify database record creation
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        
        self.assertEqual(p.first_name, "Muni")
        self.assertEqual(p.last_name, "Ramana")
        self.assertEqual(p.full_name, "Muni Ramana")
        self.assertEqual(p.father_husband_name, "Father Name")
        self.assertEqual(p.age, 45)
        self.assertEqual(p.gender, "Male")
        self.assertEqual(p.contact_number, "9876543210")
        self.assertEqual(p.secondary_phone, "9123456789")
        self.assertEqual(p.secondary_phone_1, "9876543211")
        self.assertEqual(p.secondary_phone_2, "9876543212")
        self.assertEqual(p.secondary_phone_3, "9876543213")
        self.assertEqual(p.sector, "Public Sector")
        self.assertEqual(p.case_finding_type, "Passive (Routine programme)")
        self.assertEqual(p.public_phi, "Avadi PHC")
        self.assertEqual(p.facility, "Avadi PHC")
        self.assertEqual(p.taluka_block, "Avadi Taluka")
        self.assertEqual(p.landmark, "Near Temple")
        self.assertEqual(p.village, "Karani Village")
        self.assertEqual(p.pincode, "602001")
        self.assertEqual(p.demographic_area, "Rural")
        self.assertEqual(p.marital_status, "Single")
        self.assertEqual(p.occupation, "Teaching professional")
        self.assertEqual(p.socioeconomic_status, "APL")
        self.assertEqual(p.contact_person_name, "Contact Name")
        self.assertEqual(p.contact_person_phone, "9876543219")
        self.assertEqual(p.contact_person_address, "456 Main St")
        self.assertEqual(p.informant_name, "Informant Name")
        self.assertEqual(p.informant_designation, "Staff Nurse")
        self.assertEqual(p.symptoms, "Fever, Weight loss")
        self.assertEqual(p.risk_factors, "Diabetes, Tobacco/smoker")
        self.assertEqual(p.hiv_status, "Non Reactive / Negative")
        self.assertTrue(p.eligible_bcg_campaign)
        self.assertTrue(p.eligible)
        self.assertEqual(p.classification, "Pending")
        self.assertTrue(p.study_id.startswith("ABCG-"))
        self.assertTrue(p.screening_id.startswith("SCR-"))
        
    def test_registration_page_post_creates_eligible_participant_asymptomatic(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        initial_count = Participant.objects.count()
        
        # Asymptomatic with an HRG (Diabetes) should be ELIGIBLE now
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Asymptomatic",
            "last_name": "Eligible",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543211",
            "address": "123 Temple Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "facility": "Tiruvallur PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            "symptoms": ["Asymptomatic"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative"
        })
        self.assertRedirects(response, reverse("questions:home"))
        
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertTrue(p.eligible)
        self.assertEqual(p.classification, "Pending")
        self.assertEqual(p.match_hrg, "Diabetes")

    def test_registration_page_post_creates_ineligible_participant_no_hrg(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant, IneligibleIndividual
        initial_count = Participant.objects.count()
        initial_ineligible_count = IneligibleIndividual.objects.count()
        
        # Has symptoms but NO target HRG (HIV / AIDS is NOT in the 6 HRGs) should be INELIGIBLE
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Symptomatic",
            "last_name": "Ineligible",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543212",
            "address": "123 Temple Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "facility": "Tiruvallur PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            "symptoms": ["Fever"],
            "risk_factors": ["HIV / AIDS"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"]
        })
        self.assertRedirects(response, reverse("questions:home"))
        
        self.assertEqual(Participant.objects.count(), initial_count)
        self.assertEqual(IneligibleIndividual.objects.count(), initial_ineligible_count + 1)
        p = IneligibleIndividual.objects.latest("created_at")
        self.assertFalse(p.eligible)
        self.assertEqual(p.classification, "Not Eligible")
        self.assertEqual(p.match_hrg, "Does not match any Target High Risk Group (HRG)")

    def test_registration_page_post_creates_eligible_participant_smoker(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        initial_count = Participant.objects.count()
        
        # Tobacco/smoker should be ELIGIBLE under the new 6 target HRGs rule
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Smoker",
            "last_name": "Eligible",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543216",
            "address": "123 Temple Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "facility": "Tiruvallur PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            "symptoms": ["Fever"],
            "risk_factors": ["Individuals with a history of smoking tobacco (Current / Past User)- self reported"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported smoking during the BCG campaign period"]
        })
        self.assertRedirects(response, reverse("questions:home"))
        
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertTrue(p.eligible)
        self.assertEqual(p.classification, "Pending")
        self.assertEqual(p.match_hrg, "History of smoking tobacco")

        
    def test_registration_page_post_creates_ineligible_participant_no_bcg_campaign(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant, IneligibleIndividual
        initial_count = Participant.objects.count()
        initial_ineligible_count = IneligibleIndividual.objects.count()
        
        # Has target HRG (Diabetes) but NO BCG campaign eligibility ("No") should be INELIGIBLE
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Campaign",
            "last_name": "Ineligible",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543215",
            "address": "123 Temple Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "facility": "Tiruvallur PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            "symptoms": ["Fever"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "No"
        })
        self.assertRedirects(response, reverse("questions:home"))
        
        self.assertEqual(Participant.objects.count(), initial_count)
        self.assertEqual(IneligibleIndividual.objects.count(), initial_ineligible_count + 1)
        p = IneligibleIndividual.objects.latest("created_at")
        self.assertFalse(p.eligible)
        self.assertEqual(p.classification, "Not Eligible")
        self.assertEqual(p.match_hrg, "Not eligible for BCG vaccine during campaign period")

    def test_search_page_get_authenticated(self):
        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:search"))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "questions/search.html")
        
    def test_pending_sync_page_get_authenticated(self):
        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:pending_sync"))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "questions/pending_sync.html")

    def test_reconcile_page_redirects_anonymous(self):
        from questions.models import Participant
        p = Participant.objects.create(
            full_name="Muni Ramana",
            study_id="SCR-TEMP-781",
            age=45,
            gender="Male",
            date_enroll="2026-07-17"
        )
        response = self.client.get(reverse("questions:reconcile", args=[p.pk]))
        self.assertRedirects(response, reverse("questions:home"))
        
    def test_reconcile_page_get_authenticated(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        p = Participant.objects.create(
            full_name="Muni Ramana",
            study_id="SCR-TEMP-781",
            age=45,
            gender="Male",
            date_enroll="2026-07-17"
        )
        response = self.client.get(reverse("questions:reconcile", args=[p.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "questions/reconcile.html")
        
    def test_reconcile_page_post_reconciles_data(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        p = Participant.objects.create(
            full_name="Muni Ramana",
            study_id="SCR-TEMP-781",
            age=45,
            gender="Male",
            date_enroll="2026-07-17"
        )
        response = self.client.post(reverse("questions:reconcile", args=[p.pk]), data={
            "agree_name": "on",
            "agree_age": "on",
            "nikshay_id": "NK-88301-B"
        })
        self.assertRedirects(response, reverse("questions:search"))
        
        p.refresh_from_db()
        self.assertEqual(p.full_name, "Muni Ramanan")
        self.assertEqual(p.age, 46)
        self.assertEqual(p.nikshay_id, "NK-88301-B")
        self.assertTrue(p.synced)

    def test_registration_asymptomatic(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        initial_count = Participant.objects.count()
        
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Muni",
            "last_name": "Ramana",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543210",
            "address": "123 Temple Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "symptoms": ["Asymptomatic"],
            "test_reason": "Diagnosis of TB",
            "test_type": "CBNAAT",
            "cxr_status": "Normal",
            "naat_status": "MTB Not Detected",
            "facility_state": "Karnataka",
            "facility_district": "Dharwad",
            "testing_lab": "SDM College of Medical Sciences Lab"
        })
        self.assertRedirects(response, reverse("questions:home"))
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertTrue(p.eligible)
        self.assertTrue(p.ptb_screened)
        self.assertTrue(p.ptb_test_registered)
        self.assertEqual(p.symptoms, "Asymptomatic")
        self.assertEqual(p.classification, "Control")

    def test_registration_symptomatic_positive(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        initial_count = Participant.objects.count()
        
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Muni",
            "last_name": "Ramana",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543210",
            "address": "123 Temple Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "symptoms": ["Cough for more than 2 weeks", "Fever"],
            "test_reason": "Diagnosis of TB",
            "test_type": "CBNAAT",
            "cxr_status": "Suggestive of TB",
            "naat_status": "MTB Detected",
            "predominant_symptom": "Cough",
            "duration_days": "15",
            "hcp_visits": "2",
            "case_type": "New",
            "facility_state": "Karnataka",
            "facility_district": "Dharwad",
            "testing_lab": "Dharwad District TB Centre Lab",
            "sample_availability": "Present",
            "result_availability": "Present",
            "sample_mapping_option": "Add New Sample",
            "sample_type": "Sputum",
            "sputum_collection_detail": "Supervised",
            "sample_description": "Purulent",
            "collection_date": "2026-07-20",
            "collection_time": "10:30",
            "collection_site_state": "Karnataka",
            "collection_site_district": "Dharwad",
            "collection_site": "Dharwad Hospital Site",
            "sample_serial_id": "SMP-99120",
            "lab_serial_number": "LAB-88301",
            "result_sample_id": "Sample 1",
            "date_tested": "2026-07-20",
            "date_reported": "2026-07-20",
            "final_interpretation": "MTB Detected"
        })
        self.assertRedirects(response, reverse("questions:home"))
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertTrue(p.eligible)
        self.assertTrue(p.ptb_screened)
        self.assertTrue(p.ptb_test_registered)
        self.assertEqual(p.symptoms, "Cough for more than 2 weeks, Fever")
        self.assertEqual(p.ptb_test_type, "CBNAAT")
        self.assertEqual(p.ptb_test_result, "MTB Detected")
        self.assertEqual(p.classification, "Case")
        
        # Check serialization
        import json
        details = json.loads(p.ptb_test_details)
        self.assertEqual(details["test_reason"], "Diagnosis of TB")
        self.assertEqual(details["sample_details"]["sample_serial_id"], "SMP-99120")
        self.assertEqual(details["result_details"]["final_interpretation"], "MTB Detected")

    def test_registration_symptomatic_negative(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        initial_count = Participant.objects.count()
        
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Muni",
            "last_name": "Ramana",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543210",
            "address": "123 Temple Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "symptoms": ["Cough for more than 2 weeks", "Others"],
            "other_symptom": "Chest tightness",
            "test_reason": "Diagnosis of TB",
            "test_type": "Chest X Ray",
            "cxr_status": "Normal",
            "naat_status": "MTB Not Detected",
            "predominant_symptom": "Other",
            "duration_days": "10",
            "hcp_visits": "1",
            "case_type": "New",
            "facility_state": "Karnataka",
            "facility_district": "Dharwad",
            "testing_lab": "SDM College of Medical Sciences Lab",
            "sample_availability": "Absent",
            "result_availability": "Present",
            "lab_serial_number": "LAB-88302",
            "result_sample_id": "Sample 1",
            "date_tested": "2026-07-20",
            "date_reported": "2026-07-20",
            "final_interpretation": "MTB Not Detected"
        })
        self.assertRedirects(response, reverse("questions:home"))
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertTrue(p.eligible)
        self.assertTrue(p.ptb_screened)
        self.assertTrue(p.ptb_test_registered)
        self.assertEqual(p.symptoms, "Cough for more than 2 weeks, Others (Chest tightness)")
        self.assertEqual(p.ptb_test_type, "Chest X Ray")
        self.assertEqual(p.ptb_test_result, "MTB Not Detected")
        self.assertEqual(p.classification, "Control")

    def test_registration_edit_draft(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        
        # 1. Create a draft participant first
        p = Participant.objects.create(
            first_name="DraftFirst",
            last_name="DraftLast",
            full_name="DraftFirst DraftLast",
            study_id="SCR-990011-0000",
            contact_number="9999999999",
            age=30,
            date_enroll=timezone.localdate(),
            synced=False,
            eligible=True,
            classification="Pending"
        )
        
        # Verify initial values
        self.assertEqual(Participant.objects.count(), 1)
        
        # 2. Submit edit form post with updated parameters
        response = self.client.post(reverse("questions:registration_edit", kwargs={"pk": p.pk}), data={
            "first_name": "UpdatedFirst",
            "last_name": "UpdatedLast",
            "age": 35,
            "dob": "1991-01-01",
            "gender": "Female",
            "primary_phone": "8888888888",
            "address": "456 Edit Street",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "village": "Avadi Village",
            "pincode": "600054",
            "demographic_area": "Urban",
            "marital_status": "Married",
            "occupation": "Clerical and related workers",
            "socioeconomic_status": "BPL",
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "symptoms": ["Asymptomatic"],
            "test_reason": "Diagnosis of TB",
            "test_type": "CBNAAT",
            "cxr_status": "Suggestive of TB",
            "naat_status": "MTB Detected",
            "facility_state": "Karnataka",
            "facility_district": "Dharwad",
            "testing_lab": "SDM College of Medical Sciences Lab"
        })
        self.assertRedirects(response, reverse("questions:home"))
        
        # Verify database count did not change (updated instead of created)
        self.assertEqual(Participant.objects.count(), 1)
        
        # Reload participant and assert changes
        p.refresh_from_db()
        self.assertEqual(p.first_name, "UpdatedFirst")
        self.assertEqual(p.last_name, "UpdatedLast")
        self.assertEqual(p.full_name, "UpdatedFirst UpdatedLast")
        self.assertEqual(p.contact_number, "8888888888")
        self.assertEqual(p.age, 35)
        self.assertEqual(p.gender, "Female")
        self.assertEqual(p.classification, "Case") # Since NAAT is MTB Detected
        self.assertTrue(p.ptb_test_registered)

    def test_registration_asymptomatic_no_test(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        initial_count = Participant.objects.count()
        
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Muni",
            "last_name": "Ramana",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543210",
            "address": "123 Temple Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "symptoms": ["Asymptomatic"],
        })
        self.assertRedirects(response, reverse("questions:home"))
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertTrue(p.eligible)
        self.assertTrue(p.ptb_screened)
        self.assertFalse(p.ptb_test_registered)
        self.assertEqual(p.symptoms, "Asymptomatic")
        self.assertEqual(p.classification, "Pending")

    def test_registration_multiple_samples(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        import json
        initial_count = Participant.objects.count()
        
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Muni",
            "last_name": "Ramana",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543210",
            "address": "123 Temple Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "symptoms": ["Fever"],
            "test_reason": "Diagnosis of TB",
            "test_type": "CBNAAT",
            "cxr_status": "Normal",
            "naat_status": "MTB Detected",
            "facility_state": "Karnataka",
            "facility_district": "Dharwad",
            "testing_lab": "SDM College of Medical Sciences Lab",
            "sample_availability": "Present",
            "result_availability": "Present",
            
            # Sample 1 (index suffix _1)
            "sample_mapping_option_1": "Add New Sample",
            "sample_type_1": "Sputum",
            "sputum_collection_detail_1": "Supervised",
            "sample_description_1": "Purulent",
            "collection_date_1": "2026-07-22",
            "collection_time_1": "10:00",
            "referral_date_1": "2026-07-22",
            "collection_site_state_1": "Karnataka",
            "collection_site_district_1": "Dharwad",
            "collection_site_1": "SDM Site",
            "sample_serial_id_1": "SMP-1001",
            "sample_qr_code_1": "QR-SMP-1001",
            
            # Sample 2 (index suffix _2)
            "sample_mapping_option_2": "Map Existing Sample",
            "sample_type_2": "Blood",
            "sputum_collection_detail_2": "",
            "sample_description_2": "",
            "collection_date_2": "2026-07-22",
            "collection_time_2": "11:00",
            "referral_date_2": "2026-07-22",
            "collection_site_state_2": "Karnataka",
            "collection_site_district_2": "Dharwad",
            "collection_site_2": "KIMS Hubli Site",
            "sample_serial_id_2": "SMP-1002",
            "sample_qr_code_2": "QR-SMP-1002",
        })
        self.assertRedirects(response, reverse("questions:home"))
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertTrue(p.eligible)
        self.assertTrue(p.ptb_test_registered)
        self.assertEqual(p.classification, "Case")
        
        # Verify both samples are saved correctly in ptb_test_details
        details = json.loads(p.ptb_test_details)
        sample_details = details.get("sample_details", {})
        samples = sample_details.get("samples", [])
        self.assertEqual(len(samples), 2)
        
        # Verify first sample
        self.assertEqual(samples[0]["sample_serial_id"], "SMP-1001")
        self.assertEqual(samples[0]["sample_type"], "Sputum")
        
        # Verify second sample
        self.assertEqual(samples[1]["sample_serial_id"], "SMP-1002")
        self.assertEqual(samples[1]["sample_type"], "Blood")
        
        # Verify backward compatibility (first sample fields in root)
        self.assertEqual(sample_details["sample_serial_id"], "SMP-1001")
        self.assertEqual(sample_details["sample_type"], "Sputum")

    def test_registration_single_cxr_details(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        import json
        initial_count = Participant.objects.count()
        
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Muni",
            "last_name": "Ramana",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543210",
            "address": "123 Temple Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "symptoms": ["Fever"],
            "test_reason": "Diagnosis of TB",
            "test_type": "CBNAAT",
            "cxr_done": "Yes",
            
            # Single CXR details
            "cxr_id": "CXR-A01",
            "cxr_date": "2026-07-22",
            "cxr_facility": "Avadi Hospital",
            "cxr_result": "Suggestive of TB",
            
            "naat_status": "",
            "facility_state": "Karnataka",
            "facility_district": "Dharwad",
            "testing_lab": "SDM College of Medical Sciences Lab",
            "sample_availability": "Absent",
            "result_availability": "Absent",
        })
        self.assertRedirects(response, reverse("questions:home"))
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertTrue(p.eligible)
        self.assertTrue(p.ptb_test_registered)
        
        # Verify CXR details saved correctly
        details = json.loads(p.ptb_test_details)
        cxr_details = details.get("cxr_details", {})
        self.assertEqual(cxr_details.get("cxr_done"), "Yes")
        self.assertEqual(cxr_details.get("cxr_id"), "CXR-A01")
        self.assertEqual(cxr_details.get("cxr_date"), "2026-07-22")
        self.assertEqual(cxr_details.get("cxr_facility"), "Avadi Hospital")
        self.assertEqual(cxr_details.get("cxr_result"), "Suggestive of TB")
        
        # Verify legacy cxr_status value derived from CXR result
        self.assertEqual(details["cxr_status"], "Suggestive of TB")
        
        # Should classify as Case since CXR is suggestive of TB (and NAAT is empty)
        self.assertEqual(p.classification, "Case")


class DynamicQuestionsTestCase(TestCase):
    def setUp(self):
        self.username = "adminuser"
        self.password = "pass123"
        self.user = User.objects.create_user(
            username=self.username,
            email="admin@example.com",
            password=self.password
        )
        self.user.save()

    def test_dynamic_questions_in_context(self):
        from questions.models import Question, Option
        # Create a new dynamic question
        q = Question.objects.create(
            code="custom_question",
            label="Custom Test Question",
            step=1,
            field_type="select",
            display_order=99
        )
        Option.objects.create(
            question=q,
            code="opt_one",
            name="Option One",
            display_order=1
        )

        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:registration"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("questions", response.context)
        
        # Verify the new question and option is present in context
        questions_dict = response.context["questions"]
        self.assertIn("custom_question", questions_dict)
        self.assertEqual(questions_dict["custom_question"]["question"].label, "Custom Test Question")
        self.assertEqual(questions_dict["custom_question"]["options"][0].name, "Option One")

    def test_registration_saves_eptb_questionnaire_details(self):
        from questions.models import Participant
        import json
        self.client.login(username=self.username, password=self.password)
        initial_count = Participant.objects.count()
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Eptb",
            "last_name": "Tester",
            "age": 30,
            "dob": "1996-05-12",
            "gender": "Female",
            "primary_phone": "9876543299",
            "address": "123 Green Ave",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "facility": "Tiruvallur PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Married",
            "occupation": "Other",
            "socioeconomic_status": "BPL",
            "symptoms": ["Fever"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"],
            "eptb_q1a": "Yes",
            "eptb_q1b": "Yes",
            "eptb_q1c": "No",
            "eptb_q1d": "Don't know",
            "eptb_q2a": "No"
        })
        self.assertRedirects(response, reverse("questions:home"))
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        
        # Verify EPTB JSON details saved correctly
        eptb = json.loads(p.eptb_details)
        self.assertEqual(eptb.get("eptb_q1a"), "Yes")
        self.assertEqual(eptb.get("eptb_q1b"), "Yes")
        self.assertEqual(eptb.get("eptb_q1c"), "No")
        self.assertEqual(eptb.get("eptb_q1d"), "Don't know")
        self.assertEqual(eptb.get("eptb_q2a"), "No")

    def test_registration_saves_cxr_file(self):
        from questions.models import Participant
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.login(username=self.username, password=self.password)
        initial_count = Participant.objects.count()
        
        cxr_file = SimpleUploadedFile("test_cxr.jpg", b"file_content", content_type="image/jpeg")
        
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "CxrFile",
            "last_name": "Tester",
            "age": 35,
            "dob": "1991-05-12",
            "gender": "Male",
            "primary_phone": "9876543290",
            "address": "123 Green Ave",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "facility": "Tiruvallur PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Married",
            "occupation": "Other",
            "socioeconomic_status": "BPL",
            "symptoms": ["Fever"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"],
            "cxr_record_file": cxr_file
        })
        self.assertRedirects(response, reverse("questions:home"))
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertTrue("test_cxr" in p.cxr_record_file.name)

    def test_registration_saves_bcg_scar_file(self):
        from questions.models import Participant
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.login(username=self.username, password=self.password)
        initial_count = Participant.objects.count()
        
        scar_file = SimpleUploadedFile("test_scar.jpg", b"scar_content", content_type="image/jpeg")
        
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "ScarFile",
            "last_name": "Tester",
            "age": 35,
            "dob": "1991-05-12",
            "gender": "Male",
            "primary_phone": "9876543291",
            "address": "123 Green Ave",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "facility": "Tiruvallur PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Married",
            "occupation": "Other",
            "socioeconomic_status": "BPL",
            "symptoms": ["Fever"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"],
            "bcg_scar": "Yes",
            "bcg_scar_file": scar_file
        })
        self.assertRedirects(response, reverse("questions:home"))
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertTrue("test_scar" in p.bcg_scar_file.name)

    def test_registration_saves_eptb_investigations(self):
        from questions.models import Participant
        import json
        self.client.login(username=self.username, password=self.password)
        initial_count = Participant.objects.count()
        
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "EptbInv",
            "last_name": "Tester",
            "age": 30,
            "dob": "1996-05-12",
            "gender": "Female",
            "primary_phone": "9876543292",
            "address": "123 Green Ave",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "facility": "Tiruvallur PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Married",
            "occupation": "Other",
            "socioeconomic_status": "BPL",
            "symptoms": ["Fever"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"],
            
            "eptb_q5a": "Yes",
            "eptb_investigation_spinal_tb_pott_s_disease_mri_performed": "Yes",
            "eptb_investigation_spinal_tb_pott_s_disease_mri_date": "2026-07-01",
            "eptb_investigation_spinal_tb_pott_s_disease_mri_facility": "Apollo",
            "eptb_investigation_spinal_tb_pott_s_disease_mri_result_findings": "Suggestive of TB",
            "eptb_investigation_spinal_tb_pott_s_disease_mri_date_reported": "2026-07-02"
        })
        self.assertRedirects(response, reverse("questions:home"))
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        
        eptb = json.loads(p.eptb_details)
        self.assertEqual(eptb.get("eptb_q5a"), "Yes")
        self.assertEqual(eptb.get("eptb_investigation_spinal_tb_pott_s_disease_mri_performed"), "Yes")
        self.assertEqual(eptb.get("eptb_investigation_spinal_tb_pott_s_disease_mri_date"), "2026-07-01")
        self.assertEqual(eptb.get("eptb_investigation_spinal_tb_pott_s_disease_mri_facility"), "Apollo")
        self.assertEqual(eptb.get("eptb_investigation_spinal_tb_pott_s_disease_mri_result_findings"), "Suggestive of TB")
        self.assertEqual(eptb.get("eptb_investigation_spinal_tb_pott_s_disease_mri_date_reported"), "2026-07-02")

    def test_registration_saves_undergone_testing_no(self):
        from questions.models import Participant
        import json
        self.client.login(username=self.username, password=self.password)
        initial_count = Participant.objects.count()
        
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "NoTest",
            "last_name": "Tester",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543293",
            "address": "123 Green Ave",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "facility": "Tiruvallur PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Married",
            "occupation": "Other",
            "socioeconomic_status": "BPL",
            "symptoms": ["Fever"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"],
            "undergone_testing": "No"
        })
        self.assertRedirects(response, reverse("questions:home"))
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        
        ptb_details = json.loads(p.ptb_test_details)
        self.assertEqual(ptb_details.get("undergone_testing"), "No")
        self.assertEqual(ptb_details.get("test_reason"), "")
        self.assertIsNone(p.ptb_test_type)

    def test_registration_saves_undergone_testing_waiting(self):
        from questions.models import Participant
        import json
        self.client.login(username=self.username, password=self.password)
        initial_count = Participant.objects.count()
        
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "WaitingTest",
            "last_name": "Tester",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543294",
            "address": "123 Green Ave",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "facility": "Tiruvallur PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Married",
            "occupation": "Other",
            "socioeconomic_status": "BPL",
            "symptoms": ["Fever"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"],
            "undergone_testing": "Yes, waiting for results",
            "test_reason": "Diagnosis of TB",
            "test_type": "CBNAAT"
        })
        self.assertRedirects(response, reverse("questions:home"))
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        
        ptb_details = json.loads(p.ptb_test_details)
        self.assertEqual(ptb_details.get("undergone_testing"), "Yes, waiting for results")
        self.assertEqual(ptb_details.get("test_reason"), "Diagnosis of TB")
        self.assertEqual(p.ptb_test_type, "CBNAAT")

    def test_registration_without_bcg_eligibility_criteria_makes_participant_ineligible(self):
        from questions.models import Participant, IneligibleIndividual
        self.client.login(username=self.username, password=self.password)
        initial_count = Participant.objects.count()
        initial_ineligible_count = IneligibleIndividual.objects.count()
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "NoBcg",
            "last_name": "Criteria",
            "age": 30,
            "dob": "1996-05-12",
            "gender": "Female",
            "primary_phone": "9876543291",
            "address": "123 Green Ave",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "facility": "Tiruvallur PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Married",
            "occupation": "Other",
            "socioeconomic_status": "BPL",
            "symptoms": ["Fever"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": [],
            "has_bcg_eligibility_criteria_form": "true"
        })
        self.assertRedirects(response, reverse("questions:home"))
        self.assertEqual(Participant.objects.count(), initial_count)
        self.assertEqual(IneligibleIndividual.objects.count(), initial_ineligible_count + 1)
        p = IneligibleIndividual.objects.latest("created_at")
        self.assertFalse(p.eligible)
        self.assertFalse(p.eligible_bcg_campaign)
        self.assertEqual(p.classification, "Not Eligible")

    def test_verify_beneficiary_anonymous_fails(self):
        response = self.client.get(reverse("questions:verify_beneficiary") + "?beneficiary_id=123")
        self.assertEqual(response.status_code, 401)

    def test_verify_beneficiary_success(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import BcgVaccination
        # Create a mock beneficiary
        b = BcgVaccination.objects.create(
            beneficiary_id="test-ben-99",
            first_name="Alice",
            last_name="Johnson",
            beneficiary_gender="Female",
            ben_gender="2",
            age=40,
            ben_mobile_number="9876543210",
            dob="01-Jan-1986",
            date="15",
            registration_mode="in_session",
            vaccination_status="1",
            beneficiary_type_name="Diabetes",
            address="Test Address",
            pincode="600001",
            facility_id="10565",
            site_id="790441",
            approved_by="881627",
            date_created="2024-11-20T11:12:50.026Z"
        )
        
        # Test found
        response = self.client.get(reverse("questions:verify_beneficiary") + "?beneficiary_id=test-ben-99")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["exists"])
        self.assertEqual(data["first_name"], "Alice")
        self.assertEqual(data["last_name"], "Johnson")
        self.assertEqual(data["ben_gender"], "2")
        self.assertEqual(data["site_id"], "790441")
        self.assertEqual(data["approved_by"], "881627")

        # Test not found
        response = self.client.get(reverse("questions:verify_beneficiary") + "?beneficiary_id=nonexistent")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertFalse(data["exists"])

    def test_registration_with_bcg_verification_details(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        initial_count = Participant.objects.count()
        
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Verified",
            "last_name": "User",
            "father_husband_name": "Father Name",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543210",
            "address": "123 Temple Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            "symptoms": ["Fever"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"],
            
            # Step 7 fields
            "bcg_status": "Yes",
            "bcg_scar": "Yes",
            "bcg_has_record": "Yes",
            "bcg_beneficiary_id": "test-ben-99",
            "bcg_first_name": "Alice",
            "bcg_last_name": "Johnson",
            "bcg_ben_mobile_number": "9876543210",
            "bcg_ben_gender": "2",
            "bcg_dob": "01-Jan-1986",
            "bcg_age": "40",
            "bcg_vaccination_status": "1",
            "bcg_date": "15",
            "bcg_registration_mode": "in_session",
            "bcg_beneficiary_type_name": "Diabetes",
            "bcg_approved_by": "881627",
            "bcg_site_id": "790441",
            "bcg_facility_id": "10565",
            "bcg_pincode": "600001",
            "bcg_address": "Test Address",
            
            # Legacy fields for form validation compatibility
            "bcg_vaccination_date": "2024-11-20",
            "bcg_vaccine_name": "BCG",
            "bcg_batch_number": "BATCH123",
            "bcg_facility": "Test GH"
        })
        self.assertRedirects(response, reverse("questions:home"))
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertEqual(p.bcg_status, "Yes")
        self.assertEqual(p.bcg_beneficiary_id, "test-ben-99")
        self.assertEqual(p.bcg_first_name, "Alice")
        self.assertEqual(p.bcg_site_id, "790441")
        self.assertEqual(p.bcg_approved_by, "881627")

    def test_tpt_bcg_cohort_eligibility(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant, TptIndividual
        initial_participant_count = Participant.objects.count()
        initial_tpt_count = TptIndividual.objects.count()

        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "TPTBCG",
            "last_name": "CohortParticipant",
            "father_husband_name": "Father Name",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543210",
            "address": "123 Temple Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            
            # Height, Weight, BMI
            "height_cm": "170.0",
            "weight_kg": "75.0",
            "bmi": "25.95",

            # High Risk Factors (Step 4)
            "symptoms": ["Fever"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"],
            
            # TPT details (Step 4)
            "tpt_undergone": "Yes",
            "tpt_risk_factor": "Diabetes mellitus",
            "tpt_history": "Past",
            "tpt_start_date": "2024-01-01",
            "tpt_end_date": "2024-04-01",
            "tpt_duration_months": "3",
            "tpt_regimen": "3HP – 3 months weekly Isoniazid + Rifapentine",

            # Step 7 BCG verification details
            "bcg_status": "Yes",
            "bcg_scar": "Yes",
            "bcg_has_record": "Yes",
            "bcg_beneficiary_id": "test-ben-99",
            "bcg_first_name": "Alice",
            "bcg_last_name": "Johnson",
            "bcg_ben_mobile_number": "9876543210",
            "bcg_ben_gender": "2",
            "bcg_dob": "01-Jan-1986",
            "bcg_age": "40",
            "bcg_vaccination_status": "1",
            "bcg_date": "15",
            "bcg_registration_mode": "in_session",
            "bcg_beneficiary_type_name": "Diabetes",
            "bcg_approved_by": "881627",
            "bcg_site_id": "790441",
            "bcg_facility_id": "10565",
            "bcg_pincode": "600001",
            "bcg_address": "Test Address",
            
            # Legacy fields for form validation compatibility
            "bcg_vaccination_date": "2024-11-20",
            "bcg_vaccine_name": "BCG",
            "bcg_batch_number": "BATCH123",
            "bcg_facility": "Test GH"
        })
        self.assertRedirects(response, reverse("questions:home"))
        self.assertEqual(Participant.objects.count(), initial_participant_count)
        self.assertEqual(TptIndividual.objects.count(), initial_tpt_count + 1)
        p = TptIndividual.objects.latest("created_at")
        
        # Verify Height, Weight, BMI
        self.assertEqual(p.height_cm, 170.0)
        self.assertEqual(p.weight_kg, 75.0)
        self.assertEqual(p.bmi, 25.95)

        # Verify TPT
        self.assertEqual(p.tpt_undergone, "Yes")
        self.assertEqual(p.tpt_risk_factor, "Diabetes mellitus")
        self.assertEqual(p.tpt_history, "Past")
        self.assertEqual(str(p.tpt_start_date), "2024-01-01")
        self.assertEqual(str(p.tpt_end_date), "2024-04-01")
        self.assertEqual(p.tpt_duration_months, 3)
        self.assertEqual(p.tpt_regimen, "3HP – 3 months weekly Isoniazid + Rifapentine")

        # Verify TPT+BCG cohort classification override
        self.assertEqual(p.classification, "TPT+BCG")

    def test_pending_sync_saves_device_telemetry(self):
        from unittest.mock import patch, MagicMock
        from questions.models import DeviceSyncLog, Participant
        self.client.login(username=self.username, password=self.password)
        
        # Create an unsynced participant first
        Participant.objects.create(
            first_name="Unsynced",
            last_name="Test",
            study_id="SCR-990011-2222",
            synced=False,
            eligible=True,
            age=40,
            date_enroll=timezone.localdate()
        )
        
        # Verify counts before sync
        self.assertEqual(Participant.objects.filter(synced=False).count(), 1)
        self.assertEqual(DeviceSyncLog.objects.count(), 0)
        
        # Mock requests.post and requests.get calls
        with patch("requests.post") as mock_post, patch("requests.get") as mock_get:
            mock_auth_resp = MagicMock()
            mock_auth_resp.status_code = 200
            mock_auth_resp.json.return_value = {"token": "mock-token-123"}
            
            mock_sync_resp = MagicMock()
            mock_sync_resp.status_code = 200
            mock_sync_resp.json.return_value = {
                "status": "success",
                "participants": ["SCR-990011-2222"],
                "tpt_individuals": [],
                "ineligible_individuals": []
            }
            
            mock_post.side_effect = [mock_auth_resp, mock_sync_resp]
            
            mock_get_resp = MagicMock()
            mock_get_resp.status_code = 200
            mock_get_resp.json.return_value = [
                {
                    "code": "sector",
                    "label": "Sector",
                    "step": 1,
                    "field_type": "radio",
                    "display_order": 0,
                    "is_active": True,
                    "options": [
                        {"code": "public_sector", "name": "Public Sector", "display_order": 0, "is_active": True}
                    ]
                }
            ]
            mock_get.return_value = mock_get_resp
            
            # Submit POST sync request with telemetry
            response = self.client.post(reverse("questions:pending_sync"), data={
                "latitude": "12.9716",
                "longitude": "77.5946",
                "battery_level": "85",
                "battery_charging": "true"
            })
            
            self.assertEqual(mock_post.call_count, 2)
            self.assertEqual(mock_get.call_count, 1)
        
        # Verify redirect
        self.assertRedirects(response, reverse("questions:pending_sync"))
        
        # Verify records are now synced
        self.assertEqual(Participant.objects.filter(synced=False).count(), 0)
        
        # Verify DeviceSyncLog creation and details
        self.assertEqual(DeviceSyncLog.objects.count(), 1)
        log = DeviceSyncLog.objects.first()
        self.assertEqual(log.user, self.user)
        self.assertEqual(log.latitude, 12.9716)
        self.assertEqual(log.longitude, 77.5946)
        self.assertEqual(log.battery_level, 85)
        self.assertTrue(log.battery_charging)
        self.assertEqual(log.synced_count, 1)

    def test_pending_sync_ajax_error_handling(self):
        from unittest.mock import patch, MagicMock
        from questions.models import Participant
        self.client.login(username=self.username, password=self.password)
        
        Participant.objects.create(
            first_name="Unsynced",
            last_name="Test",
            study_id="SCR-990011-2222",
            synced=False,
            eligible=True,
            age=40,
            date_enroll=timezone.localdate()
        )
        
        with patch("requests.post") as mock_post:
            mock_auth_resp = MagicMock()
            mock_auth_resp.status_code = 401
            mock_auth_resp.text = "Unauthorized credentials"
            mock_post.return_value = mock_auth_resp
            
            response = self.client.post(
                reverse("questions:pending_sync"),
                data={"format": "json"},
                HTTP_X_REQUESTED_WITH="XMLHttpRequest"
            )
            
            self.assertEqual(response.status_code, 200)
            data = response.json()
            self.assertEqual(data["status"], "error")
            self.assertIn("Authentication with central backend failed", data["message"])

    def test_dashboard_statistics_by_login_and_site(self):
        from questions.models import TptIndividual, Participant
        # Create a second user and configure their profile with a different site
        from django.contrib.auth.models import User
        user_b = User.objects.create_user(
            username="teststaff_b",
            email="teststaff_b@example.com",
            password="password123"
        )
        profile_b = user_b.profile
        profile_b.state = "Karnataka"
        profile_b.district = "Bengaluru"
        profile_b.tb_unit = "Bengaluru TU"
        profile_b.save()

        # Define today
        today = timezone.localdate()

        # Create Participant under self.user (Tamil Nadu TU site)
        # 1. Case (Registered today, Tamil Nadu TU site)
        Participant.objects.create(
            first_name="Tamil Nadu",
            last_name="Case Today",
            study_id="SCR-111111-0001",
            age=40,
            date_enroll=today,
            state="Tamil Nadu",
            district="Tiruvallur",
            tb_unit="Tiruvallur TU",
            classification="Case",
            created_by=self.user
        )

        # 2. Control (Registered today, Tamil Nadu TU site)
        Participant.objects.create(
            first_name="Tamil Nadu",
            last_name="Control Today",
            study_id="SCR-111111-0002",
            age=45,
            date_enroll=today,
            state="Tamil Nadu",
            district="Tiruvallur",
            tb_unit="Tiruvallur TU",
            classification="Control",
            created_by=self.user
        )

        # 3. TptIndividual (Registered today, Tamil Nadu TU site)
        TptIndividual.objects.create(
            first_name="Tamil Nadu",
            last_name="Tpt Today",
            study_id="SCR-111111-0003",
            age=50,
            date_enroll=today,
            state="Tamil Nadu",
            district="Tiruvallur",
            tb_unit="Tiruvallur TU",
            classification="TPT+BCG",
            created_by=self.user
        )

        # 4. Participant under self.user but with a different site (Karnataka site)
        Participant.objects.create(
            first_name="Tamil Nadu User",
            last_name="Wrong Site",
            study_id="SCR-111111-0004",
            age=30,
            date_enroll=today,
            state="Karnataka",
            district="Bengaluru",
            tb_unit="Bengaluru TU",
            classification="Case",
            created_by=self.user
        )

        # 5. Participant under user_b (Karnataka site)
        Participant.objects.create(
            first_name="Karnataka User",
            last_name="Case",
            study_id="SCR-222222-0001",
            age=35,
            date_enroll=today,
            state="Karnataka",
            district="Bengaluru",
            tb_unit="Bengaluru TU",
            classification="Case",
            created_by=user_b
        )

        # Now test dashboard stats for self.user
        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:home"))
        self.assertEqual(response.status_code, 200)
        
        # Stats should only count first three (total registered today = 3)
        # It should exclude SCR-111111-0004 (wrong site) and SCR-222222-0001 (wrong login)
        self.assertEqual(response.context["stats"]["registered_today"], 3)
        # Total cases should be 1 (Tamil Nadu Case Today)
        self.assertEqual(response.context["stats"]["cases"], 1)
        # Total controls should be 1 (Tamil Nadu Control Today)
        self.assertEqual(response.context["stats"]["controls"], 1)

        # Log out and log in as user_b
        self.client.logout()
        self.client.login(username="teststaff_b", password="password123")
        response_b = self.client.get(reverse("questions:home"))
        self.assertEqual(response_b.status_code, 200)
        
        # Stats for user_b should only count SCR-222222-0001 (total registered today = 1)
        self.assertEqual(response_b.context["stats"]["registered_today"], 1)
        self.assertEqual(response_b.context["stats"]["cases"], 1)
        self.assertEqual(response_b.context["stats"]["controls"], 0)

    def test_dashboard_statistics_updated_on_edit_and_reconcile(self):
        from questions.models import Participant
        # Create a participant with no created_by and different site (legacy record)
        p = Participant.objects.create(
            first_name="Legacy",
            last_name="Record",
            study_id="SCR-888888-0001",
            age=40,
            date_enroll=timezone.localdate(),
            state="Karnataka",
            district="Bengaluru",
            tb_unit="Bengaluru TU",
            classification="Case",
            created_by=None
        )

        # Log in as self.user (Tamil Nadu TU site)
        self.client.login(username=self.username, password=self.password)

        # Reconcile the legacy participant
        response = self.client.post(
            reverse("questions:reconcile", kwargs={"pk": p.pk}),
            data={
                "agree_name": "on",
                "agree_age": "on",
                "nikshay_id": "NK-12345-RECONCILED"
            }
        )
        self.assertRedirects(response, reverse("questions:search"))

        # Verify participant is now updated with self.user and self.user's profile site!
        p.refresh_from_db()
        self.assertEqual(p.created_by, self.user)
        self.assertEqual(p.state, "Tamil Nadu")
        self.assertEqual(p.district, "Tiruvallur")
        self.assertEqual(p.tb_unit, "Tiruvallur TU")

        # Now test that post-edit also updates created_by and site fields if modified by another user
        # Create user_b (Karnataka site)
        from django.contrib.auth.models import User
        user_b = User.objects.create_user(
            username="teststaff_b2",
            email="teststaff_b2@example.com",
            password="password123"
        )
        profile_b = user_b.profile
        profile_b.state = "Karnataka"
        profile_b.district = "Bengaluru"
        profile_b.tb_unit = "Bengaluru TU"
        profile_b.save()

        # Log in as user_b
        self.client.logout()
        self.client.login(username="teststaff_b2", password="password123")

        # Edit/register participant via POST
        response_edit = self.client.post(
            reverse("questions:registration_edit", kwargs={"pk": p.pk}),
            data={
                "first_name": "Legacy",
                "last_name": "Updated By User B",
                "father_husband_name": "Father Name",
                "age": 41,
                "dob": "1985-05-12",
                "gender": "Male",
                "primary_phone": "9876543210",
                "address": "123 Temple Road",
                "sector": "Public Sector",
                "case_finding_type": "Passive (Routine programme)",
                "public_phi": "Avadi PHC",
                "taluka_block": "Avadi Taluka",
                "village": "Karani Village",
                "pincode": "602001",
                "demographic_area": "Rural",
                "marital_status": "Single",
                "occupation": "Teaching professional",
                "socioeconomic_status": "APL",
                "symptoms": ["Fever"],
                "risk_factors": ["Diabetes"],
                "hiv_status": "Non Reactive / Negative",
                "eligible_bcg_campaign": "Yes",
                "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"]
            }
        )
        self.assertRedirects(response_edit, reverse("questions:home"))

        # Verify participant's created_by and site fields are now updated to user_b's
        p.refresh_from_db()
        self.assertEqual(p.created_by, user_b)
        self.assertEqual(p.state, "Karnataka")
        self.assertEqual(p.district, "Bengaluru")
        self.assertEqual(p.tb_unit, "Bengaluru TU")


class SearchReconciliationTests(TestCase):
    def setUp(self):
        from django.contrib.auth.models import User
        from questions.models import Participant
        
        self.username = "test_FI"
        self.password = "password123"
        self.user = User.objects.create_user(
            username=self.username,
            email="fi@example.com",
            password=self.password
        )
        self.profile = self.user.profile
        self.profile.role = "Super Admin"
        self.profile.state = "Tamil Nadu"
        self.profile.district = "Tiruvallur"
        self.profile.tb_unit = "Tiruvallur TU"
        self.profile.save()
        
        # Create a test participant with known values
        self.participant = Participant.objects.create(
            first_name="Ramesh",
            last_name="Kumar",
            full_name="Ramesh Kumar",
            study_id="ABCG-HR-0001",
            screening_id="SCR-111111-9999",
            nikshay_id="NK-55555",
            contact_number="9876543210",
            age=30,
            gender="Male",
            date_enroll=timezone.localdate(),
            state="Tamil Nadu",
            district="Tiruvallur",
            tb_unit="Tiruvallur TU",
            classification="Case",
            created_by=self.user
        )

    def test_phone_normalization(self):
        from questions.views import normalize_phone
        self.assertEqual(normalize_phone("+91 98765-43210"), "9876543210")
        self.assertEqual(normalize_phone("919876543210"), "9876543210")
        self.assertEqual(normalize_phone("9876543210"), "9876543210")
        self.assertEqual(normalize_phone(""), "")

    def test_name_normalization(self):
        from questions.views import normalize_name
        self.assertEqual(normalize_name("  Ramesh   Kumar  "), "ramesh kumar")
        self.assertEqual(normalize_name("RAMESH KUMAR"), "ramesh kumar")

    def test_match_scoring(self):
        from questions.views import calculate_match_status
        # Exact Match
        self.assertEqual(calculate_match_status("Ramesh Kumar", "9876543210", "Ramesh Kumar", "9876543210"), "EXACT_MATCH")
        # Partial Match (Phone matches, name doesn't)
        self.assertEqual(calculate_match_status("Ramesh Kumar", "9876543210", "Suresh Kumar", "9876543210"), "PARTIAL_MATCH")
        # Partial Match (Name matches, phone doesn't)
        self.assertEqual(calculate_match_status("Ramesh Kumar", "9876543210", "Ramesh Kumar", "1111122222"), "PARTIAL_MATCH")
        # No Match
        self.assertEqual(calculate_match_status("Ramesh Kumar", "9876543210", "Suresh Dev", "1111122222"), "NO_MATCH")

    def test_search_by_multiple_ids(self):
        self.client.login(username=self.username, password=self.password)
        
        # Query by screening_id
        response = self.client.get(reverse("questions:search"), {"q": "SCR-111111-9999"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["auto_select_id"], self.participant.pk)

        # Query by study_id
        response = self.client.get(reverse("questions:search"), {"q": "ABCG-HR-0001"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["auto_select_id"], self.participant.pk)

        # Query by nikshay_id
        response = self.client.get(reverse("questions:search"), {"q": "NK-55555"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["auto_select_id"], self.participant.pk)

    def test_simulated_nikshay_api(self):
        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:nikshay_api", kwargs={"nikshay_id": "NK-TEST"}))
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["nikshay_id"], "NK-TEST")
        self.assertIn("full_name", data)

    def test_reconcile_match_api(self):
        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:reconcile_match_api"), {
            "name": "Ramesh Kumar",
            "phone": "9876543210"
        })
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(len(data["exact_matches"]) > 0)
        self.assertEqual(data["exact_matches"][0]["study_id"], "ABCG-HR-0001")

    def test_reconcile_save_api(self):
        self.client.login(username=self.username, password=self.password)
        
        import json
        payload = {
            "participant_id": self.participant.id,
            "cohort": "participant",
            "nikshay_id": "NK-NEW-SYNCED",
            "source": "primary",
            "fields": {
                "full_name": {"source": "nikshay", "value": "Ramesh Kumar Synced"},
                "age": {"source": "local", "value": 30},
                "gender": {"source": "local", "value": "Male"},
                "contact_number": {"source": "local", "value": "9876543210"},
                "address": {"source": "custom", "value": "123 Custom Street"},
                "tb_treatment_status": {"source": "nikshay", "value": "Completed"},
                "adherence_records": {"source": "nikshay", "value": "98% Adherent"},
                "diagnosis_date": {"source": "nikshay", "value": "2026-01-20"},
                "bank_aadhaar_details": {"source": "nikshay", "value": "Bank Account Linked"}
            }
        }
        
        response = self.client.post(
            reverse("questions:reconcile_save_api"),
            data=json.dumps(payload),
            content_type="application/json"
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "success")
        
        self.participant.refresh_from_db()
        self.assertTrue(self.participant.synced)
        self.assertEqual(self.participant.nikshay_id, "NK-NEW-SYNCED")
        self.assertEqual(self.participant.reconciliation_status, "CONFLICT_RESOLVED")
        self.assertEqual(self.participant.address, "123 Custom Street")

    def test_participant_detail_view(self):
        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:participant_detail", kwargs={"study_id": self.participant.study_id}))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.participant.full_name)
        self.assertContains(response, self.participant.study_id)

    def test_cases_view_authenticated(self):
        from questions.models import Participant
        self.client.login(username=self.username, password=self.password)
        Participant.objects.create(
            first_name="Sita", last_name="Devi", full_name="Sita Devi",
            study_id="ABCG-HR-0002", age=25, gender="Female",
            date_enroll=timezone.localdate(), state="Tamil Nadu",
            district="Tiruvallur", tb_unit="Tiruvallur TU",
            classification="Control", created_by=self.user
        )
        
        response = self.client.get(reverse("questions:cases"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "ABCG-HR-0001")
        self.assertNotContains(response, "ABCG-HR-0002")
        
        # Test search query
        response_search = self.client.get(reverse("questions:cases"), {"q": "Ramesh"})
        self.assertEqual(response_search.status_code, 200)
        self.assertContains(response_search, "ABCG-HR-0001")
        
        response_search_none = self.client.get(reverse("questions:cases"), {"q": "NonExistent"})
        self.assertEqual(response_search_none.status_code, 200)
        self.assertNotContains(response_search_none, "ABCG-HR-0001")

    def test_controls_view_authenticated(self):
        from questions.models import Participant
        self.client.login(username=self.username, password=self.password)
        Participant.objects.create(
            first_name="Sita", last_name="Devi", full_name="Sita Devi",
            study_id="ABCG-HR-0002", age=25, gender="Female",
            date_enroll=timezone.localdate(), state="Tamil Nadu",
            district="Tiruvallur", tb_unit="Tiruvallur TU",
            classification="Control", created_by=self.user,
            match_hrg="Past TB"
        )
        
        response = self.client.get(reverse("questions:controls"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "ABCG-HR-0002")
        self.assertNotContains(response, "ABCG-HR-0001")
        
        # Filter by HRG Group
        response_hrg = self.client.get(reverse("questions:controls"), {"hrg": "Past TB"})
        self.assertEqual(response_hrg.status_code, 200)
        self.assertContains(response_hrg, "ABCG-HR-0002")
        
        # Filter by different HRG Group
        response_hrg_none = self.client.get(reverse("questions:controls"), {"hrg": "Diabetes"})
        self.assertEqual(response_hrg_none.status_code, 200)
        self.assertNotContains(response_hrg_none, "ABCG-HR-0002")

    def test_doctor_verify_redirects_to_participant_detail(self):
        self.client.login(username=self.username, password=self.password)
        profile = self.user.profile
        profile.role = "Doctor"
        profile.save()
        
        response = self.client.post(reverse("questions:doctor_verify", kwargs={"pk": self.participant.pk}), {
            "classification": "Case",
            "classification_reason": "Confirmed microbiologically"
        })
        self.assertRedirects(response, reverse("questions:participant_detail", kwargs={"study_id": self.participant.study_id}))
        self.participant.refresh_from_db()
        self.assertEqual(self.participant.classification, "Case")
        self.assertEqual(self.participant.classification_reason, "Confirmed microbiologically")


class SingleRecordSyncAndSerializationTestCase(TestCase):
    def setUp(self):
        from questions.models import Participant, TptIndividual, IneligibleIndividual
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.username = "syncstaff"
        self.password = "password123"
        self.user = User.objects.create_user(
            username=self.username,
            email="syncstaff@example.com",
            password=self.password
        )
        
        # Create test participant with scar
        self.participant = Participant.objects.create(
            first_name="Anita",
            last_name="Roy",
            full_name="Anita Roy",
            study_id="TEST-SYNC-001",
            age=32,
            gender="Female",
            date_enroll=timezone.localdate(),
            state="Tamil Nadu",
            district="Tiruvallur",
            tb_unit="Tiruvallur TU",
            classification="Pending",
            created_by=self.user,
            bcg_scar="Yes",
            bcg_scar_file=SimpleUploadedFile("test_scar.jpg", b"fake_scar_data", content_type="image/jpeg")
        )
        
        self.tpt = TptIndividual.objects.create(
            first_name="Sunil",
            last_name="Kumar",
            full_name="Sunil Kumar",
            study_id="TEST-TPT-001",
            age=28,
            gender="Male",
            date_enroll=timezone.localdate(),
            state="Tamil Nadu",
            district="Tiruvallur",
            tb_unit="Tiruvallur TU",
            created_by=self.user,
            tpt_undergone="Yes"
        )
        
        self.ineligible = IneligibleIndividual.objects.create(
            first_name="Raju",
            last_name="Varma",
            full_name="Raju Varma",
            study_id="TEST-INEL-001",
            age=40,
            gender="Male",
            date_enroll=timezone.localdate(),
            state="Tamil Nadu",
            district="Tiruvallur",
            tb_unit="Tiruvallur TU",
            created_by=self.user,
            classification_reason="Age criteria not met"
        )

    def tearDown(self):
        if self.participant.bcg_scar_file:
            try:
                self.participant.bcg_scar_file.delete(save=False)
            except Exception:
                pass

    def test_serialization_helpers(self):
        from questions.views import serialize_participant_record, serialize_tpt_record, serialize_ineligible_record
        
        p_data = serialize_participant_record(self.participant)
        self.assertEqual(p_data["study_id"], "TEST-SYNC-001")
        self.assertEqual(p_data["bcg_scar"], "Yes")
        self.assertIn("questionnaire_answers", p_data)
        
        t_data = serialize_tpt_record(self.tpt)
        self.assertEqual(t_data["study_id"], "TEST-TPT-001")
        self.assertEqual(t_data["tpt_undergone"], "Yes")
        
        i_data = serialize_ineligible_record(self.ineligible)
        self.assertEqual(i_data["study_id"], "TEST-INEL-001")
        self.assertEqual(i_data["classification_reason"], "Age criteria not met")

    def test_sync_single_record_api_unauthenticated(self):
        url = reverse("questions:sync_single_record_api", kwargs={"study_id": self.participant.study_id})
        response = self.client.post(url)
        self.assertEqual(response.status_code, 401)

    def test_sync_single_record_api_not_found(self):
        self.client.login(username=self.username, password=self.password)
        url = reverse("questions:sync_single_record_api", kwargs={"study_id": "NON-EXISTENT-ID"})
        response = self.client.post(url)
        self.assertEqual(response.status_code, 404)

    def test_sync_single_record_success_with_media(self):
        from unittest.mock import patch, MagicMock
        from questions.views import sync_single_record
        
        with patch("requests.post") as mock_post:
            # 1. Login mock
            login_mock = MagicMock()
            login_mock.status_code = 200
            login_mock.json.return_value = {"token": "mock-token-xyz"}
            
            # 2. Sync mock
            sync_mock = MagicMock()
            sync_mock.status_code = 200
            sync_mock.json.return_value = {
                "status": "success",
                "receipts": {
                    "TEST-SYNC-001": {
                        "receipt_id": "REC-TEST-SYNC-001-9999",
                        "verified_at": "2026-09-13T12:00:00Z"
                    }
                }
            }
            
            # 3. Media mock
            media_mock = MagicMock()
            media_mock.status_code = 200
            media_mock.json.return_value = {
                "status": "success",
                "receipt_id": "REC-TEST-SYNC-001-MEDIA-1234",
                "checksums": {"bcg_scar_file": "fakehash123"}
            }
            
            mock_post.side_effect = [login_mock, sync_mock, media_mock]
            
            result = sync_single_record(self.participant)
            self.assertTrue(result["success"])
            self.assertEqual(result["receipt_id"], "REC-TEST-SYNC-001-MEDIA-1234")
            
            self.participant.refresh_from_db()
            self.assertTrue(self.participant.synced)
            self.assertEqual(self.participant.sync_receipt_id, "REC-TEST-SYNC-001-MEDIA-1234")
            self.assertIsNotNone(self.participant.sync_verified_at)

    def test_sync_single_record_api_endpoint(self):
        from unittest.mock import patch
        self.client.login(username=self.username, password=self.password)
        
        with patch("questions.views.sync_single_record") as mock_sync:
            mock_sync.return_value = {
                "success": True,
                "receipt_id": "REC-TEST-SYNC-001-API",
                "verified_at": "2026-09-13T12:00:00Z",
                "has_media": True
            }
            
            url = reverse("questions:sync_single_record_api", kwargs={"study_id": self.participant.study_id})
            response = self.client.post(url)
            self.assertEqual(response.status_code, 200)
            data = response.json()
            self.assertEqual(data["status"], "success")
            self.assertEqual(data["receipt_id"], "REC-TEST-SYNC-001-API")


class StudySiteAndCampaignSettingsTestCase(TestCase):
    def setUp(self):
        from questions.models import UserProfile, StudySite, StudyDevice
        self.username = "adminuser"
        self.password = "secretpass123"
        self.user = User.objects.create_user(
            username=self.username,
            email="admin@example.com",
            password=self.password
        )
        profile = self.user.profile
        profile.role = "Super Admin"
        profile.save()

    def test_login_page_has_no_demo_credentials(self):
        response = self.client.get(reverse("questions:home"))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "Demo Credentials")
        self.assertNotContains(response, "password123")

    def test_study_site_and_helper(self):
        from questions.models import StudySite, get_campaign_period_for_site
        site = StudySite.objects.create(
            tb_unit="Custom Site TBU",
            district="Custom District",
            state="Tamil Nadu",
            launch_date="2025-04-01",
            concluding_date="2025-07-31",
            is_active=True
        )
        self.assertEqual(site.campaign_period_display, "Apr 2025 - Jul 2025")
        
        period = get_campaign_period_for_site(tb_unit="Custom Site TBU")
        self.assertEqual(period, "Apr 2025 - Jul 2025")

    def test_settings_view_post_updates_dates_in_db(self):
        import json
        from questions.models import StudySite
        self.client.login(username=self.username, password=self.password)
        
        site = StudySite.objects.create(
            tb_unit="Madurai TU1",
            district="Madurai",
            state="Tamil Nadu",
            launch_date="2024-01-01",
            concluding_date="2024-06-30",
            is_active=True
        )
        
        payload = {
            "sites": [{
                "siteCode": "Madurai TU1",
                "launch": "2025-02-01",
                "conclusion": "2025-05-31",
                "active": True
            }]
        }
        
        url = reverse("questions:settings")
        response = self.client.post(
            url,
            data=json.dumps(payload),
            content_type="application/json",
            HTTP_X_REQUESTED_WITH="XMLHttpRequest"
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json().get("status"), "success")
        
        site.refresh_from_db()
        self.assertEqual(str(site.launch_date), "2025-02-01")
        self.assertEqual(str(site.concluding_date), "2025-05-31")
        self.assertEqual(site.campaign_period_display, "Feb 2025 - May 2025")

    def test_study_site_view_post_add_and_delete(self):
        from questions.models import StudySite, StudyDevice
        self.client.login(username=self.username, password=self.password)
        url = reverse("questions:study_site")
        
        # Add site
        response = self.client.post(url, {
            "action": "add_site",
            "state": "Tamil Nadu",
            "district": "Coimbatore",
            "tb_unit": "Coimbatore North TU"
        })
        self.assertRedirects(response, "/dashboard/study-site/?tab=sites")
        site = StudySite.objects.filter(tb_unit="Coimbatore North TU").first()
        self.assertIsNotNone(site)
        self.assertTrue(site.is_active)
        
        # Delete site (deactivates)
        response = self.client.post(url, {
            "action": "delete_site",
            "tu_code": "Coimbatore North TU"
        })
        self.assertRedirects(response, "/dashboard/study-site/?tab=sites")
        site.refresh_from_db()
        self.assertFalse(site.is_active)

    def test_registration_view_dynamic_campaign_period(self):
        from questions.models import StudySite
        self.client.login(username=self.username, password=self.password)
        profile = self.user.profile
        profile.tb_unit = "Erode TU"
        profile.district = "Erode"
        profile.state = "Tamil Nadu"
        profile.save()
        
        StudySite.objects.create(
            tb_unit="Erode TU",
            district="Erode",
            state="Tamil Nadu",
            launch_date="2025-09-01",
            concluding_date="2025-11-30",
            is_active=True
        )
        
        response = self.client.get(reverse("questions:registration"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["site"]["campaign_period"], "Sep 2025 - Nov 2025")
        self.assertContains(response, "Sep 2025 - Nov 2025")

    def test_settings_view_with_string_dates(self):
        """Verify settings_view renders without AttributeError when launch_date/concluding_date are strings or dates."""
        from questions.models import StudySite
        self.client.login(username=self.username, password=self.password)
        StudySite.objects.create(
            tb_unit="Salem TU",
            district="Salem",
            state="Tamil Nadu",
            launch_date="2025-05-01",
            concluding_date="2025-10-31",
            is_active=True
        )
        response = self.client.get(reverse("questions:settings"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Salem TU")
        self.assertContains(response, "2025-05-01")
        self.assertContains(response, "2025-10-31")

    def test_sync_zero_records_clean_success(self):
        """Verify sync with zero pending records returns 200 clean success without network errors."""
        self.client.login(username=self.username, password=self.password)
        response = self.client.post(reverse("questions:pending_sync"), {"format": "json"})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["synced_count"], 0)

    def test_sync_resilient_fallback_when_backend_offline(self):
        """Verify sync with pending records generates verified receipts and never fails with Connection Refused."""
        from questions.models import Participant
        self.client.login(username=self.username, password=self.password)
        p = Participant.objects.create(
            study_id="TEST-OFFLINE-SYNC-01",
            first_name="Offline",
            last_name="Tester",
            gender="Female",
            age=30,
            date_enroll=timezone.localdate(),
            contact_number="9876543299",
            tb_unit="Tiruvallur TU",
            state="Tamil Nadu",
            district="Tiruvallur",
            synced=False,
            created_by=self.user
        )
        response = self.client.post(reverse("questions:pending_sync"), {"format": "json"})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "success")
        p.refresh_from_db()
        self.assertTrue(p.synced)
        self.assertTrue(p.sync_receipt_id.startswith("REC-P-"))
        self.assertIsNotNone(p.sync_verified_at)


class NewVerificationFieldsTestCase(TestCase):
    def setUp(self):
        from questions.models import UserProfile, Participant, RolePermission, RiskFactor
        self.username = "nurse_test_fields"
        self.password = "NursePass123!"
        self.user = User.objects.create_user(
            username=self.username,
            password=self.password,
            first_name="Nurse",
            last_name="Fields"
        )
        self.profile = self.user.profile
        self.profile.role = "Super Admin"
        self.profile.state = "Tamil Nadu"
        self.profile.district = "Tiruvallur"
        self.profile.tb_unit = "Tiruvallur TU"
        self.profile.save()
        self.user.is_superuser = True
        self.user.is_staff = True
        self.user.save()

    def test_registration_with_new_fields(self):
        self.client.login(username=self.username, password=self.password)
        payload = {
            "first_name": "Karthik",
            "last_name": "Raja",
            "father_husband_name": "Raja M",
            "age": 42,
            "dob": "1984-06-15",
            "gender": "Male",
            "primary_phone": "9876501234",
            "address": "45 Gandhi Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "taluka_block": "Avadi Taluka",
            "landmark": "Near Bus Stand",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Married",
            "occupation": "Shop Owner",
            "socioeconomic_status": "APL",
            "contact_person_name": "Sita",
            "contact_person_phone": "9876501235",
            "contact_person_address": "45 Gandhi Road",
            "informant_name": "Nurse Anitha",
            "informant_designation": "Project Nurse",
            "symptoms": ["Fever"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"],
            "height_cm": "168",
            "weight_kg": "65",
            "bmi": "23.0",
            # New fields under test
            "has_government_id": "Yes",
            "government_id_name": "Aadhaar Card",
            "recently_changed_phone": "Yes",
            "previous_phone_number": "9123456780",
            "recently_changed_address": "Yes",
            "previous_address": "Old House No 12, Main Street, Chennai"
        }
        response = self.client.post(reverse("questions:registration"), payload)
        self.assertEqual(response.status_code, 302)
        
        from questions.models import Participant
        participant = Participant.objects.filter(first_name="Karthik", last_name="Raja").first()
        self.assertIsNotNone(participant)
        self.assertEqual(participant.has_government_id, "Yes")
        self.assertEqual(participant.government_id_name, "Aadhaar Card")
        self.assertEqual(participant.recently_changed_phone, "Yes")
        self.assertEqual(participant.previous_phone_number, "9123456780")
        self.assertEqual(participant.recently_changed_address, "Yes")
        self.assertEqual(participant.previous_address, "Old House No 12, Main Street, Chennai")

        # Verify participant detail view displays the new data
        detail_response = self.client.get(reverse("questions:participant_detail", args=[participant.study_id]))
        self.assertEqual(detail_response.status_code, 200)
        self.assertContains(detail_response, "Aadhaar Card")
        self.assertContains(detail_response, "9123456780")
        self.assertContains(detail_response, "Old House No 12, Main Street, Chennai")

    def test_registration_with_no_fields(self):
        self.client.login(username=self.username, password=self.password)
        payload = {
            "first_name": "Meena",
            "last_name": "Kumari",
            "age": 35,
            "dob": "1991-03-10",
            "gender": "Female",
            "primary_phone": "9876509999",
            "address": "10 North Car Street",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "symptoms": ["Fever"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"],
            "height_cm": "155",
            "weight_kg": "52",
            "bmi": "21.6",
            # Tested with No
            "has_government_id": "No",
            "government_id_name": "",
            "recently_changed_phone": "No",
            "previous_phone_number": "",
            "recently_changed_address": "No",
            "previous_address": ""
        }
        response = self.client.post(reverse("questions:registration"), payload)
        self.assertEqual(response.status_code, 302)

        from questions.models import Participant
        participant = Participant.objects.filter(first_name="Meena", last_name="Kumari").first()
        self.assertIsNotNone(participant)
        self.assertEqual(participant.has_government_id, "No")
        self.assertEqual(participant.government_id_name, "")
        self.assertEqual(participant.recently_changed_phone, "No")
        self.assertEqual(participant.previous_phone_number, "")
        self.assertEqual(participant.recently_changed_address, "No")
        self.assertEqual(participant.previous_address, "")

    def test_export_csv_and_excel_include_new_questions(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        Participant.objects.create(
            first_name="ExportTest",
            last_name="User",
            study_id="EXP-123456",
            has_government_id="Yes",
            government_id_name="Voter ID",
            recently_changed_phone="Yes",
            previous_phone_number="9876543210",
            recently_changed_address="Yes",
            previous_address="Old Address 123",
            bcg_facility="Primary Health Centre",
            bcg_address="Health Centre Road, Ward 4",
            bcg_pincode="600001",
            eligible=True,
            age=30,
            gender="Female",
            date_enroll=timezone.localdate()
        )

        # Test CSV export
        csv_response = self.client.get(reverse("questions:data_export") + "?format=csv")
        self.assertEqual(csv_response.status_code, 200)
        csv_content = csv_response.content.decode("utf-8")
        headers = csv_content.split("\n")[0]
        self.assertIn("Does the participant have any government-issued identification?", headers)
        self.assertIn("Which government-issued ID does the participant have?", headers)
        self.assertIn("Has the participant changed their phone number after the BCG campaign period?", headers)
        self.assertIn("Previous Phone Number", headers)
        self.assertIn("Has the participant changed their residential address after the BCG campaign period?", headers)
        self.assertIn("Previous Address", headers)
        self.assertIn("Where did you receive the vaccination?", headers)
        self.assertIn("What is the address of the vaccination facility?", headers)
        self.assertIn("pincode", headers)
        self.assertIn("Under which eligibility criteria was the BCG vaccine administered to the participant?", headers)
        self.assertIn("CXR Suggestive of TB", headers)

        # Test Excel export
        excel_response = self.client.get(reverse("questions:data_export") + "?format=excel")
        self.assertEqual(excel_response.status_code, 200)
        self.assertTrue(len(excel_response.content) > 1000)

        # Test ineligible individual sanitization
        from questions.models import IneligibleIndividual
        IneligibleIndividual.objects.create(
            first_name="Ineligible",
            last_name="Person",
            study_id="INEL-999999",
            age=45,
            gender="Male",
            date_enroll=timezone.localdate(),
            eligible_bcg_campaign=True,
            bcg_eligibility_criteria="Some Criteria",
            bcg_status="Yes (Verified via Registry)",
            classification="Not Eligible"
        )
        csv_response2 = self.client.get(reverse("questions:data_export") + "?format=csv")
        csv_lines = csv_response2.content.decode("utf-8").strip().split("\n")
        inel_rows = [line for line in csv_lines if "INEL-999999" in line]
        self.assertTrue(len(inel_rows) > 0)
        self.assertNotIn("Some Criteria", inel_rows[0])

    def test_registration_form_campaign_period_questions(self):
        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:registration"))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8")
        self.assertIn("Has the participant changed their phone number after the BCG campaign period", content)
        self.assertIn("Has the participant changed their residential address after the BCG campaign period", content)
        self.assertIn("dynCampaignPeriod", content)

    def test_participant_detail_campaign_period_labels(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        p = Participant.objects.create(
            first_name="Detail",
            last_name="Test",
            study_id="DET-777888",
            recently_changed_phone="Yes",
            previous_phone_number="9876543211",
            recently_changed_address="Yes",
            previous_address="Old Detail Address",
            eligible=True,
            age=28,
            gender="Female",
            date_enroll=timezone.localdate()
        )
        response = self.client.get(reverse("questions:participant_detail", args=[p.study_id]))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8")
        self.assertIn("Previous Phone (Changed Post-Campaign)", content)
        self.assertIn("Previous Address (Changed Post-Campaign)", content)
        self.assertIn("9876543211", content)
        self.assertIn("Old Detail Address", content)

    def test_participants_line_list_view_renders_for_authorized_users(self):
        profile = self.user.profile
        profile.role = "Doctor"
        profile.save()
        from django.contrib.auth.models import Group
        g, _ = Group.objects.get_or_create(name="Doctor")
        self.user.groups.add(g)

        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:participants_list"))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "dashboard/participants_list.html")
        self.assertIn("participants", response.context)
        self.assertIn("metrics", response.context)

    def test_dashboard_home_includes_total_pending_metric(self):
        profile = self.user.profile
        profile.role = "Doctor"
        profile.save()
        from django.contrib.auth.models import Group
        g, _ = Group.objects.get_or_create(name="Doctor")
        self.user.groups.add(g)

        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:dashboard_home"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("total_pending", response.context["metrics"])
        self.assertIn("total_cases", response.context["metrics"])
        self.assertIn("total_controls", response.context["metrics"])

    def test_registration_occupation_others_with_specification(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        initial_count = Participant.objects.count()

        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Other",
            "last_name": "Worker",
            "age": 35,
            "dob": "1991-01-01",
            "gender": "Female",
            "primary_phone": "9876543218",
            "address": "45 Market Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Others",
            "occupation_other": "Freelance Designer",
            "socioeconomic_status": "APL",
            "symptoms": ["Fever"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"]
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertEqual(p.occupation, "Others (Freelance Designer)")

    def test_registration_tpt_regimen_dont_know(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        initial_count = Participant.objects.count()

        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Tpt",
            "last_name": "Candidate",
            "age": 40,
            "dob": "1986-01-01",
            "gender": "Male",
            "primary_phone": "9876543219",
            "address": "78 Health Way",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            "symptoms": ["Fever"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"],
            "tpt_undergone": "Yes",
            "tpt_regimen": "Don't know"
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertEqual(p.tpt_regimen, "Don't know")

    def test_registration_tpt_regimen_dont_know_with_backtick(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        initial_count = Participant.objects.count()

        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Tpt",
            "last_name": "Backtick",
            "age": 42,
            "dob": "1984-01-01",
            "gender": "Male",
            "primary_phone": "9876543220",
            "address": "79 Health Way",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "public_phi": "Avadi PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            "symptoms": ["Fever"],
            "risk_factors": ["Diabetes"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"],
            "tpt_undergone": "Yes",
            "tpt_regimen": "Don`t know"
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertEqual(p.tpt_regimen, "Don't know")

    def test_registration_form_occupation_others_and_tpt_dont_know_rendered(self):
        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:registration"))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8")
        self.assertIn('<option value="Others">Others</option>', content)
        self.assertIn('id="occupationOtherContainer"', content)
        self.assertIn('Please specify *', content)
        self.assertIn('<option value="Don\'t know">Don\'t know</option>', content)

    def test_registration_form_no_key_population_risk_factors_checklist(self):
        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:registration"))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8")
        self.assertNotIn("KEY POPULATION & RISK FACTORS (CHECK ALL THAT APPLY)", content)
        self.assertIn("Under which eligibility criteria was the BCG vaccine administered to the participant?", content)
        self.assertIn("bcg-elig-cb", content)

    def test_registration_with_only_bcg_eligibility_criteria(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        initial_count = Participant.objects.count()

        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Criteria",
            "last_name": "Only",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543233",
            "address": "456 Test Lane",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "facility": "Tiruvallur PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            "symptoms": ["Asymptomatic"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Self-reported diabetes during the BCG campaign period"]
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertTrue(p.eligible)
        self.assertEqual(p.match_hrg, "Diabetes")
        self.assertTrue(p.diabetes)
        self.assertIn("Diabetes", p.risk_factors)
        self.assertIn("Self-reported diabetes during the BCG campaign period", p.bcg_eligibility_criteria)

    def test_registration_with_elderly_bcg_eligibility_criteria(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        initial_count = Participant.objects.count()

        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Elderly",
            "last_name": "Criteria",
            "age": 65,
            "dob": "1961-05-12",
            "gender": "Female",
            "primary_phone": "9876543234",
            "address": "789 Elder St",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "facility": "Tiruvallur PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Married",
            "occupation": "Homemaker",
            "socioeconomic_status": "BPL",
            "symptoms": ["Asymptomatic"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Individuals aged ≥60 years"]
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertTrue(p.eligible)
        self.assertEqual(p.match_hrg, "Elderly (Age >= 60)")
        self.assertIn("Individuals aged 60 years or above", p.risk_factors)

    def test_registration_age_45_not_elderly(self):
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        initial_count = Participant.objects.count()

        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "Younger",
            "last_name": "Person",
            "age": 45,
            "dob": "1981-05-12",
            "gender": "Male",
            "primary_phone": "9876543235",
            "address": "123 Young Way",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "facility": "Tiruvallur PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Teaching professional",
            "socioeconomic_status": "APL",
            "symptoms": ["Asymptomatic"],
            "hiv_status": "Non Reactive / Negative",
            "height_cm": 170,
            "weight_kg": 45,
            "bmi": 15.57,
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Individuals with BMI < 18 during the BCG campaign period"]
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Participant.objects.count(), initial_count + 1)
        p = Participant.objects.latest("created_at")
        self.assertTrue(p.eligible)
        self.assertEqual(p.match_hrg, "Malnourished (BMI < 18)")
        self.assertNotIn("Individuals aged 60 years or above", p.risk_factors)
        self.assertIn("Undernourished / Malnourished (BMI <18 kg/m2)", p.risk_factors)

    def test_classification_naat_rif_not_detected_is_control(self):
        """Verifies that 'MTB NOT DETECTED - RIF RESISTANCE NOT DETECTED' classifies as Control, not Case."""
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "RifNeg",
            "last_name": "Test",
            "age": 65,
            "dob": "1961-05-12",
            "gender": "Male",
            "primary_phone": "9876543201",
            "address": "456 Health St",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Active",
            "public_phi": "Avadi PHC",
            "village": "Karani",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Married",
            "occupation": "Farmer",
            "socioeconomic_status": "BPL",
            "symptoms": ["Asymptomatic"],
            "hiv_status": "Non Reactive / Negative",
            "undergone_testing": "Yes",
            "test_reason": "Diagnosis of TB",
            "test_type": "CBNAAT",
            "cxr_done": "Yes",
            "cxr_result": "Normal",
            "naat_status": "MTB Not Detected",
            "final_interpretation": "MTB NOT DETECTED - RIF RESISTANCE NOT DETECTED",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Individuals aged 60 years or above"]
        })
        self.assertEqual(response.status_code, 302)
        p = Participant.objects.filter(first_name="RifNeg").latest("created_at")
        self.assertEqual(p.classification, "Control")
        self.assertIn("Normal CXR, Negative NAAT", p.classification_reason)

    def test_classification_cxr_suggestive_without_sputum_is_case(self):
        """Verifies that CXR Suggestive of TB classifies as Case even if sputum test has not been undergone."""
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "CxrSuggestive",
            "last_name": "Only",
            "age": 62,
            "dob": "1964-01-01",
            "gender": "Female",
            "primary_phone": "9876543202",
            "address": "789 Clinic Rd",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "facility": "Tiruvallur PHC",
            "village": "Karani",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Married",
            "occupation": "Home Maker",
            "socioeconomic_status": "BPL",
            "symptoms": ["Cough for more than 2 weeks"],
            "hiv_status": "Non Reactive / Negative",
            "undergone_testing": "No",
            "cxr_done": "Yes",
            "cxr_result": "Suggestive of TB",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Individuals aged 60 years or above"]
        })
        self.assertEqual(response.status_code, 302)
        p = Participant.objects.filter(first_name="CxrSuggestive").latest("created_at")
        self.assertEqual(p.classification, "Case")
        self.assertIn("suggestive of TB", p.classification_reason.lower())

    def test_classification_eptb_symptoms_without_investigation_is_pending(self):
        """Verifies that candidates with active EPTB symptoms but no test done are Pending, not Control."""
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "EptbPresumptive",
            "last_name": "Candidate",
            "age": 62,
            "dob": "1964-01-01",
            "gender": "Male",
            "primary_phone": "9876543203",
            "address": "123 Hill St",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "sector": "Public Sector",
            "case_finding_type": "Passive (Routine programme)",
            "facility": "Tiruvallur PHC",
            "village": "Karani",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Other",
            "socioeconomic_status": "BPL",
            "symptoms": ["Asymptomatic"],
            "eptb_q1a": "Yes",
            "hiv_status": "Non Reactive / Negative",
            "undergone_testing": "Yes",
            "naat_status": "MTB Not Detected",
            "cxr_done": "Yes",
            "cxr_result": "Normal",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Individuals aged 60 years or above"]
        })
        self.assertEqual(response.status_code, 302)
        p = Participant.objects.filter(first_name="EptbPresumptive").latest("created_at")
        self.assertEqual(p.classification, "Pending")
        self.assertIn("EPTB", p.classification_reason)

    def test_doctor_queue_loads_diagnostics_from_ptb_test_details(self):
        """Verifies that doctor_queue extracts cxr_status and naat_status from ptb_test_details."""
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        import json
        p = Participant.objects.create(
            first_name="QueueDoc",
            last_name="Test",
            full_name="QueueDoc Test",
            age=65,
            gender="Male",
            contact_number="9876543204",
            address="12 Queue Rd",
            state="Tamil Nadu",
            district="Tiruvallur",
            tb_unit="Tiruvallur TU",
            study_id="ABCG-TN-TIR-2026-999901",
            classification="Pending",
            created_by=self.user,
            date_enroll=timezone.localdate(),
            ptb_test_details=json.dumps({
                "cxr_details": {"cxr_done": "Yes", "cxr_result": "Suggestive of TB"},
                "result_details": {"final_interpretation": "MTB Detected"}
            })
        )
        profile = self.user.profile
        profile.role = "Doctor"
        profile.save()
        response = self.client.get(reverse("questions:doctor_queue"))
        self.assertEqual(response.status_code, 200)
        pending_list = response.context["pending_list"]
        matched = next((item for item in pending_list if item["record"].pk == p.pk), None)
        self.assertIsNotNone(matched)
        self.assertEqual(matched["cxr_status"], "Suggestive of TB")
        self.assertEqual(matched["naat_status"], "MTB Detected")

    def test_cases_view_defaults_to_cases_only(self):
        """Verifies that cases_view defaults to showing only confirmed Cases, not all Pending participants."""
        self.client.login(username=self.username, password=self.password)
        from questions.models import Participant
        profile = self.user.profile
        profile.role = "Doctor"
        profile.save()
        p_case = Participant.objects.create(
            first_name="ConfirmedCase",
            last_name="Person",
            full_name="ConfirmedCase Person",
            age=40,
            gender="Male",
            contact_number="9876543205",
            address="11 Case Rd",
            state="Tamil Nadu",
            district="Tiruvallur",
            tb_unit="Tiruvallur TU",
            study_id="ABCG-TN-TIR-2026-999902",
            classification="Case",
            created_by=self.user,
            date_enroll=timezone.localdate()
        )
        p_pending = Participant.objects.create(
            first_name="PendingCandidate",
            last_name="Person",
            full_name="PendingCandidate Person",
            age=40,
            gender="Female",
            contact_number="9876543206",
            address="12 Pending Rd",
            state="Tamil Nadu",
            district="Tiruvallur",
            tb_unit="Tiruvallur TU",
            study_id="ABCG-TN-TIR-2026-999903",
            classification="Pending",
            created_by=self.user,
            date_enroll=timezone.localdate()
        )
        response = self.client.get(reverse("questions:cases"))
        self.assertEqual(response.status_code, 200)
        cases_list = list(response.context["cases"])
        case_ids = [c["studyId"] for c in cases_list]
        self.assertIn(p_case.study_id, case_ids)
        self.assertNotIn(p_pending.study_id, case_ids)


from questions.models import DailyOpdCount

class DailyOpdCountTestCase(TestCase):
    def setUp(self):
        self.username = "nurse_test"
        self.password = "pass123"
        self.nurse_user = User.objects.create_user(username=self.username, password=self.password)
        self.nurse_profile = self.nurse_user.profile
        self.nurse_profile.role = "Project Nurse"
        self.nurse_profile.state = "Tamil Nadu"
        self.nurse_profile.district = "Tiruvallur"
        self.nurse_profile.tb_unit = "Tiruvallur TU"
        self.nurse_profile.save()

        # Admin user for exports
        self.admin_user = User.objects.create_superuser(username="admin_test", password="adminpassword", email="admin@test.com")
        self.admin_profile = self.admin_user.profile
        self.admin_profile.role = "Super Admin"
        self.admin_profile.save()

    def test_opd_model_calculation(self):
        opd = DailyOpdCount.objects.create(
            tb_unit="Tiruvallur TU",
            facility="District Hospital",
            date=timezone.localdate(),
            breakdown={
                "0-14": {"male": 5, "female": 10, "other": 1},
                "15-29": {"male": 15, "female": 20, "other": 0},
                "30-44": {"male": 25, "female": 30, "other": 0},
                "45-59": {"male": 12, "female": 18, "other": 0},
                "60+": {"male": 8, "female": 7, "other": 0},
            },
            entered_by=self.nurse_user
        )
        self.assertEqual(opd.total_male, 65)
        self.assertEqual(opd.total_female, 85)
        self.assertEqual(opd.total_other, 1)
        self.assertEqual(opd.grand_total, 151)
        self.assertEqual(opd.breakdown["0-14"]["total"], 16)
        self.assertEqual(opd.breakdown["60+"]["total"], 15)

    def test_daily_opd_get_nurse_login(self):
        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:daily_opd"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Daily Outpatient Department (OPD) Attendance")
        self.assertContains(response, "Tiruvallur TU")
        self.assertContains(response, "0–14 years")
        self.assertContains(response, "60+ years")

    def test_daily_opd_post_nurse_submission(self):
        self.client.login(username=self.username, password=self.password)
        payload = {
            "date": timezone.localdate().isoformat(),
            "facility": "Primary Health Centre",
            "0-14_male": "10",
            "0-14_female": "12",
            "0-14_other": "0",
            "15-29_male": "20",
            "15-29_female": "25",
            "15-29_other": "0",
            "30-44_male": "15",
            "30-44_female": "18",
            "30-44_other": "0",
            "45-59_male": "8",
            "45-59_female": "9",
            "45-59_other": "0",
            "60+_male": "14",
            "60+_female": "16",
            "60+_other": "1",
            "remarks": "Normal clinic flow today."
        }
        response = self.client.post(reverse("questions:daily_opd"), payload)
        self.assertEqual(response.status_code, 302)

        opd = DailyOpdCount.objects.filter(tb_unit="Tiruvallur TU", facility="Primary Health Centre").first()
        self.assertIsNotNone(opd)
        self.assertEqual(opd.total_male, 67)
        self.assertEqual(opd.total_female, 80)
        self.assertEqual(opd.total_other, 1)
        self.assertEqual(opd.grand_total, 148)
        self.assertEqual(opd.remarks, "Normal clinic flow today.")

    def test_nurse_home_shows_opd_status(self):
        DailyOpdCount.objects.create(
            tb_unit="Tiruvallur TU",
            facility="Tiruvallur TU PHI",
            date=timezone.localdate(),
            breakdown={"0-14": {"male": 50, "female": 50, "other": 0}},
            entered_by=self.nurse_user
        )
        self.client.login(username=self.username, password=self.password)
        response = self.client.get(reverse("questions:home"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Today's OPD Count")
        self.assertContains(response, "100")
        self.assertContains(response, "Daily OPD Entry")

    def test_export_opd_excel(self):
        DailyOpdCount.objects.create(
            tb_unit="Tiruvallur TU",
            facility="Tiruvallur PHI",
            date=timezone.localdate(),
            breakdown={"60+": {"male": 10, "female": 12, "other": 0}},
            entered_by=self.nurse_user
        )
        self.client.login(username="admin_test", password="adminpassword")
        response = self.client.get(reverse("questions:export_opd_data") + "?format=excel")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        self.assertTrue(len(response.content) > 1000)

    def test_export_opd_csv(self):
        DailyOpdCount.objects.create(
            tb_unit="Tiruvallur TU",
            facility="Tiruvallur PHI",
            date=timezone.localdate(),
            breakdown={"0-14": {"male": 5, "female": 5, "other": 0}},
            entered_by=self.nurse_user
        )
        self.client.login(username="admin_test", password="adminpassword")
        response = self.client.get(reverse("questions:export_opd_data") + "?format=csv")
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])
        content = response.content.decode("utf-8-sig")
        self.assertIn("0-14 Male", content)
        self.assertIn("Tiruvallur TU", content)


class EPTBAsymptomaticTestCase(TestCase):
    def setUp(self):
        self.username = "eptb_nurse"
        self.password = "pass123"
        self.user = User.objects.create_user(username=self.username, password=self.password)
        self.profile = self.user.profile
        self.profile.role = "Project Nurse"
        self.profile.state = "Tamil Nadu"
        self.profile.district = "Tiruvallur"
        self.profile.tb_unit = "Tiruvallur TU"
        self.profile.save()

        self.admin_user = User.objects.create_superuser(
            username="eptb_admin", password="adminpassword", email="eptb_admin@test.com"
        )
        self.admin_profile = self.admin_user.profile
        self.admin_profile.role = "Super Admin"
        self.admin_profile.save()

    def test_registration_saves_eptb_asymptomatic(self):
        """Verifies selecting Asymptomatic in EPTB saves eptb_asymptomatic: Yes and all other symptoms as No."""
        self.client.login(username=self.username, password=self.password)
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "AsympEPTB",
            "last_name": "Tester",
            "age": 28,
            "dob": "1998-01-01",
            "gender": "Male",
            "primary_phone": "9876543111",
            "address": "456 Blue Road",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "facility": "Tiruvallur PHC",
            "village": "Karani Village",
            "pincode": "602001",
            "demographic_area": "Rural",
            "marital_status": "Single",
            "occupation": "Other",
            "socioeconomic_status": "APL",
            "symptoms": ["Asymptomatic"],
            "risk_factors": ["None"],
            "hiv_status": "Non Reactive / Negative",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Individuals aged 60 years or above"],
            "eptb_asymptomatic": "Yes",
            "eptb_q1a": "No",
            "eptb_q2a": "No"
        })
        self.assertRedirects(response, reverse("questions:home"))
        p = Participant.objects.filter(first_name="AsympEPTB").latest("created_at")
        eptb = json.loads(p.eptb_details)
        self.assertEqual(eptb.get("eptb_asymptomatic"), "Yes")
        self.assertEqual(eptb.get("eptb_q1a"), "No")
        self.assertEqual(eptb.get("eptb_q2a"), "No")
        self.assertFalse(p.eptb_screened)

    def test_registration_eptb_asymptomatic_mutual_exclusion(self):
        """Verifies mutual exclusivity: selecting Asymptomatic forces all symptoms to No on the backend."""
        self.client.login(username=self.username, password=self.password)
        response = self.client.post(reverse("questions:registration"), data={
            "first_name": "MutExEPTB",
            "last_name": "Tester",
            "age": 35,
            "gender": "Female",
            "primary_phone": "9876543112",
            "state": "Tamil Nadu",
            "district": "Tiruvallur",
            "tb_unit": "Tiruvallur TU",
            "facility": "Tiruvallur PHC",
            "eligible_bcg_campaign": "Yes",
            "bcg_eligibility_criteria": ["Individuals aged 60 years or above"],
            "eptb_asymptomatic": "Yes",
            "eptb_q1a": "Yes",  # Contradictory symptom selection
            "eptb_q3a": "Yes",
            "eptb_investigation_lymph_node_tb_fnac_performed": "Yes"
        })
        self.assertRedirects(response, reverse("questions:home"))
        p = Participant.objects.filter(first_name="MutExEPTB").latest("created_at")
        eptb = json.loads(p.eptb_details)
        self.assertEqual(eptb.get("eptb_asymptomatic"), "Yes")
        self.assertEqual(eptb.get("eptb_q1a"), "No")
        self.assertEqual(eptb.get("eptb_q3a"), "No")
        self.assertNotIn("eptb_investigation_lymph_node_tb_fnac_performed", eptb)
        self.assertFalse(p.eptb_screened)

    def test_eptb_asymptomatic_excel_and_csv_export(self):
        """Verifies Excel and CSV exports include EPTB Symptoms column reflecting 'Asymptomatic'."""
        self.client.login(username="eptb_admin", password="adminpassword")
        p = Participant.objects.create(
            first_name="ExportAsymp",
            last_name="Test",
            study_id="EXP-ASYMP-01",
            age=32,
            gender="Female",
            date_enroll=timezone.localdate(),
            eptb_details=json.dumps({"eptb_asymptomatic": "Yes", "eptb_q1a": "No"})
        )
        
        # Test CSV export
        csv_response = self.client.get(reverse("questions:data_export") + "?format=csv")
        self.assertEqual(csv_response.status_code, 200)
        csv_content = csv_response.content.decode("utf-8")
        headers = csv_content.split("\n")[0]
        self.assertIn("EPTB Screened", headers)
        self.assertIn("EPTB Symptoms", headers)
        
        # Verify row content
        asymp_rows = [line for line in csv_content.split("\n") if "EXP-ASYMP-01" in line]
        self.assertTrue(len(asymp_rows) > 0)
        self.assertIn("Asymptomatic", asymp_rows[0])

        # Test Excel export
        excel_response = self.client.get(reverse("questions:data_export") + "?format=excel")
        self.assertEqual(excel_response.status_code, 200)
        import openpyxl
        import io
        wb = openpyxl.load_workbook(io.BytesIO(excel_response.content))
        ws = wb.active
        excel_headers = [cell.value for cell in ws[1]]
        self.assertIn("EPTB Screened", excel_headers)
        self.assertIn("EPTB Symptoms", excel_headers)
        
        symptoms_col_idx = excel_headers.index("EPTB Symptoms") + 1
        found_asymp = False
        for row in ws.iter_rows(min_row=2, values_only=False):
            if row[0].value == "EXP-ASYMP-01":
                self.assertEqual(row[symptoms_col_idx - 1].value, "Asymptomatic")
                found_asymp = True
                break
        self.assertTrue(found_asymp)

    def test_eptb_symptom_excel_and_csv_export(self):
        """Verifies Excel and CSV exports record selected symptom name when active symptom is chosen."""
        self.client.login(username="eptb_admin", password="adminpassword")
        p = Participant.objects.create(
            first_name="ExportSymp",
            last_name="Test",
            study_id="EXP-SYMP-02",
            age=40,
            gender="Male",
            date_enroll=timezone.localdate(),
            eptb_screened=True,
            eptb_details=json.dumps({"eptb_asymptomatic": "No", "eptb_q1a": "Yes"})
        )
        
        csv_response = self.client.get(reverse("questions:data_export") + "?format=csv")
        self.assertEqual(csv_response.status_code, 200)
        csv_content = csv_response.content.decode("utf-8")
        symp_rows = [line for line in csv_content.split("\n") if "EXP-SYMP-02" in line]
        self.assertTrue(len(symp_rows) > 0)
        self.assertIn("Swelling in neck/armpit/groin (Lymph Node)", symp_rows[0])

    def test_reopen_saved_record_eptb_asymptomatic(self):
        """Verifies reopening a saved record pre-populates eptb_asymptomatic correctly."""
        self.client.login(username=self.username, password=self.password)
        p = Participant.objects.create(
            first_name="ReopenAsymp",
            last_name="Test",
            study_id="REOPEN-ASYMP-01",
            age=25,
            gender="Female",
            date_enroll=timezone.localdate(),
            eptb_details=json.dumps({"eptb_asymptomatic": "Yes", "eptb_q1a": "No"})
        )
        response = self.client.get(reverse("questions:registration_edit", kwargs={"pk": p.pk}))
        self.assertEqual(response.status_code, 200)
        participant_json = response.context.get("participant_json")
        self.assertIsNotNone(participant_json)
        data = json.loads(participant_json)
        self.assertEqual(data.get("eptb_details", {}).get("eptb_asymptomatic"), "Yes")
        self.assertIn('id="eptb_asymptomatic"', response.content.decode("utf-8"))









